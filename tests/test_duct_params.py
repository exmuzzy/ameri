"""Ключевые слова покупных позиций (harness/duct_calc/params.yaml) не задевают собственные изделия эталонов.

Прототип ищет слова как подстроки: слово из строки отвода или перехода перенесёт её в расчётку без расчёта.
Собственные изделия — позиции эталонов (harness/evals/, demo/training/), у которых в expected.yaml задан тип.
"""

from pathlib import Path

import docx
import openpyxl
import pytest
import yaml

from ameri.actions import duct_params

ROOT = Path(__file__).resolve().parents[1]
HARNESS = ROOT / "harness"
EVALS = sorted(p.parent for p in [*(HARNESS / "evals").glob("*/expected.yaml"), *(ROOT / "demo" / "training").glob("*/expected.yaml")])


def spec_lines(folder: Path) -> list[str]:
    spec = next(p for p in sorted(folder.iterdir()) if p.stem == "spec")
    if spec.suffix == ".xlsx":
        book = openpyxl.load_workbook(spec, read_only=True)
        return [" ".join(str(v) for v in row if v not in (None, "")) for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)]
    if spec.suffix == ".docx":
        document = docx.Document(spec)
        return [p.text for p in document.paragraphs] + [" ".join(c.text for c in row.cells) for t in document.tables for row in t.rows]
    return spec.read_text(encoding="utf-8").splitlines()


def own_product_lines(folder: Path) -> list[str]:
    expected = yaml.safe_load((folder / "expected.yaml").read_text(encoding="utf-8"))
    lines = [line.lower() for line in spec_lines(folder)]
    found = []
    for position in expected["positions"]:
        if not position.get("type"):
            continue
        hits = [line for line in lines if position["line"].lower() in line]
        assert hits, f"{folder.name}: строки «{position['line']}» нет в спецификации"
        found += hits
    return found


def test_evals_found():
    assert {"himlab", "order-142444", "zakaz-plastik"} <= {p.name for p in EVALS}


@pytest.mark.parametrize("folder", EVALS, ids=lambda p: p.name)
def test_passthrough_keywords_miss_own_products(folder):
    keywords = [str(word).lower() for word in duct_params(HARNESS).get("passthrough_keywords") or ()]
    hits = [f"«{word}» в «{line}»" for line in own_product_lines(folder) for word in keywords if word in line]
    assert not hits, "покупное слово ловит своё изделие: " + "; ".join(hits)
