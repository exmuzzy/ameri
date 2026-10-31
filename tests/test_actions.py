from pathlib import Path

import pytest

from ameri import actions
from ameri.harness import assistant_prompt, load_access
from ameri.auth import User
from ameri.store import Attachment, Store, Usage

HARNESS = Path(__file__).resolve().parents[1] / "harness"


class FakeLlm:
    def __init__(self):
        self.messages = None

    def complete(self, messages, **_):
        self.messages = messages
        return "ответ", Usage(prompt_tokens=10, completion_tokens=2, cache_hit_tokens=8)


def test_ask_assistant_includes_history_and_files(tmp_path):
    store = Store(tmp_path / "db.sqlite3", tmp_path / "files")
    chat_id = store.create_chat("t", "anna")
    store.add_message(chat_id, author="anna", role="user", content="что тут?", files=[("spec.md", b"d=200")])
    llm = FakeLlm()
    result = actions.ask_assistant(llm, "sys", store.messages(chat_id), None, {"anna": "Анна"})
    assert result.text == "ответ"
    assert llm.messages[0] == {"role": "system", "content": "sys"}
    assert "Анна: что тут?" in llm.messages[1]["content"]
    assert "d=200" in llm.messages[1]["content"]


def test_duct_calc_unavailable_without_prototype(tmp_path):
    assert not actions.duct_calc_available(None)
    assert not actions.duct_calc_available(tmp_path)


def test_harness_files_load():
    access = load_access(HARNESS)
    manager = User("m", "M", "manager", "")
    leader = User("l", "L", "leader", "")
    admin = User("a", "A", "admin", "")
    assert access.can(manager, "action.duct_calc")
    assert not access.can(manager, "chat.view_all")
    assert not access.can(manager, "harness.edit")
    assert not access.can(manager, "feedback.review")
    assert access.can(leader, "chat.view_all")
    assert access.can(leader, "harness.edit")
    assert access.can(leader, "feedback.review")
    assert access.can(admin, "harness.edit")
    assert access.can(admin, "feedback.review")
    assert "ameri" in assistant_prompt(HARNESS)


def test_history_budget_keeps_latest(tmp_path, monkeypatch):
    monkeypatch.setattr(actions, "MAX_HISTORY_CHARS", 50)
    store = Store(tmp_path / "db.sqlite3", tmp_path / "files")
    chat_id = store.create_chat("t", "anna")
    for i in range(5):
        store.add_message(chat_id, author="anna", role="user", content=f"сообщение номер {i}")
    llm = FakeLlm()
    result = actions.ask_assistant(llm, "sys", store.messages(chat_id), None, {})
    contents = [m["content"] for m in llm.messages]
    assert result.usage.cache_hit_tokens == 8
    assert "сообщение номер 4" in contents[-1]
    assert "опущены" in contents[1]
    assert not any("номер 0" in c for c in contents)


def test_prompt_order_is_static_first(tmp_path):
    import shutil
    from ameri.harness import HarnessRepo

    harness = tmp_path / "harness"
    shutil.copytree(HARNESS, harness)
    repo = HarnessRepo(harness)
    repo.add_rule("Клапан", "Клапан считать как 2 метра прямого участка.", author="boss", source="исправление #1")
    repo.add_example("Сколько стоит отвод?", "Цену считает расчётка.", author="boss")
    prompt = assistant_prompt(harness)
    assert prompt.index("Ты — ассистент") < prompt.index("Правила, утверждённые") < prompt.index("Примеры правильных")


def test_odt_text(tmp_path):
    import zipfile

    content = (
        '<office:document-content xmlns:office="o" xmlns:text="t" xmlns:table="tb"><office:body><office:text>'
        "<text:p>Тема: Воздуховоды</text:p>"
        "<table:table><table:table-row><table:table-cell><text:p>Воздуховод ПП ф160 L1000</text:p></table:table-cell>"
        "<table:table-cell><text:p>2</text:p></table:table-cell><table:table-cell><text:p>шт</text:p></table:table-cell>"
        "</table:table-row></table:table>"
        "</office:text></office:body></office:document-content>"
    )
    path = tmp_path / "spec.odt"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("content.xml", content)
    assert actions.odt_text(path) == "Тема: Воздуховоды\nВоздуховод ПП ф160 L1000 | 2 | шт"


def _attachment(path: Path) -> Attachment:
    return Attachment(id=1, name=path.name, path=path, size=path.stat().st_size if path.exists() else 0)


def test_attachment_preview_text(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("привет мир", encoding="utf-8")
    preview = actions.attachment_preview(_attachment(path))
    assert preview.kind == "text"
    assert preview.text == "привет мир"
    assert not preview.truncated


def test_attachment_preview_markdown(tmp_path):
    path = tmp_path / "spec.md"
    path.write_text("# Заголовок\n\n| A | B |\n|---|---|\n| 1 | 2 |", encoding="utf-8")
    preview = actions.attachment_preview(_attachment(path))
    assert preview.kind == "markdown"
    assert preview.text == "# Заголовок\n\n| A | B |\n|---|---|\n| 1 | 2 |"
    assert not preview.truncated


def test_attachment_preview_image_and_pdf_by_suffix(tmp_path):
    image = tmp_path / "photo.png"
    image.write_bytes(b"\x89PNG\r\n")
    assert actions.attachment_preview(_attachment(image)).kind == "image"

    pdf = tmp_path / "doc.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    assert actions.attachment_preview(_attachment(pdf)).kind == "pdf"


def test_attachment_preview_unsupported_for_legacy_doc(tmp_path):
    path = tmp_path / "старый.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0")
    assert actions.attachment_preview(_attachment(path)).kind == "unsupported"


def test_attachment_preview_docx(tmp_path):
    from docx import Document

    path = tmp_path / "spec.docx"
    document = Document()
    document.add_paragraph("Раздел 1")
    document.add_paragraph("Воздуховод ф160")
    document.save(path)
    preview = actions.attachment_preview(_attachment(path))
    assert preview.kind == "text"
    assert preview.text == "Раздел 1\nВоздуховод ф160"


def test_attachment_preview_xlsx(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "spec.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Наименование", "Кол-во"])
    sheet.append(["Воздуховод ф160", 2])
    workbook.save(path)
    preview = actions.attachment_preview(_attachment(path))
    assert preview.kind == "table"
    assert preview.rows == [["Наименование", "Кол-во"], ["Воздуховод ф160", "2"]]
    assert not preview.truncated


def test_attachment_preview_odt(tmp_path):
    import zipfile

    content = (
        '<office:document-content xmlns:office="o" xmlns:text="t"><office:body><office:text>'
        "<text:p>Тема: Воздуховоды</text:p>"
        "</office:text></office:body></office:document-content>"
    )
    path = tmp_path / "spec.odt"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("content.xml", content)
    preview = actions.attachment_preview(_attachment(path))
    assert preview.kind == "text"
    assert preview.text == "Тема: Воздуховоды"


def test_non_ascii_key_gives_clear_message():
    import pytest
    from ameri.llm import BAD_KEY_MESSAGE, DeepSeekChat, LlmError, key_problem

    assert key_problem("sk-abc123") is None
    assert key_problem("sk-ключ") == BAD_KEY_MESSAGE
    assert key_problem("sk-abc 123") == BAD_KEY_MESSAGE
    with pytest.raises(LlmError, match="недопустимые символы"):
        DeepSeekChat(base_url="https://x", api_key="sk-вставьте ключ", model="m")


def test_spec_lines_from_tables(tmp_path):
    md = tmp_path / "s.md"
    md.write_text(
        "# Приток П2\n\n| № | Наименование | Кол-во | Ед. |\n|---|---|---|---|\n"
        "| 1 | Воздуховод ПП 400×300 δ=4 | 8 | м |\n| 2 | Отвод 90° ПП ф160 | 2 | шт |\n",
        encoding="utf-8",
    )
    assert actions.spec_lines_text(md).splitlines() == [
        "Приток П2",
        "1. Воздуховод ПП 400×300 δ=4 — 8 м; кол-во: 8",
        "2. Отвод 90° ПП ф160 — 2 шт; кол-во: 2",
    ]
    csv_file = tmp_path / "s.csv"
    csv_file.write_text("№;Наименование;Кол-во;Ед. изм.\n1;Воздуховод ПП ф160 (В1);10;мп\n2;Хомут ф160;4;шт\n", encoding="utf-8")
    assert actions.spec_lines_text(csv_file).splitlines() == [
        "1. Воздуховод ПП ф160 (В1) — 10 мп; кол-во: 10",
        "2. Хомут ф160 — 4 шт; кол-во: 4",
    ]
    assert actions.spec_lines_text(tmp_path / "x.xlsx") is None


def test_spec_lines_from_odt_keeps_system_headings():
    demo = Path(__file__).resolve().parents[1] / "demo" / "specs" / "06-ceh-pokraski-v1-p1.odt"
    lines = actions.spec_lines_text(demo).splitlines()
    assert lines[0] == "Цех покраски, объект «Восточная площадка»"
    assert "1. Воздуховод ПП ∅400 δ=4 — 14 мп; система В1; кол-во: 14" in lines
    assert "9. Решётка ПП 500×400 — 2 шт; система П1; кол-во: 2" in lines
    assert not any(line.startswith("Система") for line in lines)
    assert not any("Наименование" in line for line in lines)


def test_red_positions_from_result(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([None, 1, "Воздуховод ПП ф315\t10\tмп\tКол-во: 10 мп — отрезки по 1500 мм"])
    sheet.append([None, 2, "Переход ПП ф315\t2\tшт\tКол-во: 2 шт (проблема: Позиция 5: не указан размер d_out)"])
    sheet.append([None, 3, "Отвод 90° ПП\t3\tшт (проблема: Позиция 6: тип не распознан)"])
    path = tmp_path / "r.xlsx"
    book.save(path)
    assert actions.red_positions(path) == [
        "«Переход ПП ф315 · 2 · шт» — не указан второй диаметр (выход)",
        "«Отвод 90° ПП · 3 · шт» — тип не распознан",
    ]
    assert actions.red_positions(tmp_path / "нет.xlsx") == []


def test_assistant_prompt_includes_parse_rules():
    prompt = assistant_prompt(HARNESS)
    assert "Как расчётка понимает строки спецификаций" in prompt
    assert "Тройник без размера ответвления" in prompt
    assert prompt.index("Как расчётка понимает") < prompt.index("Примеры правильных")


# --- Q38 и Q39 поверх прототипа (прототипа в репозитории нет: проверяем на его подобии) ---

from dataclasses import dataclass as _dc, field as _field
from types import SimpleNamespace as _NS


@_dc(frozen=True)
class _Question:
    code: str
    text: str


@_dc(frozen=True)
class _Position:
    source_no: str
    source_text: str
    element_type: str | None = None
    qty: float | None = None
    length_mm: float | None = None
    size: dict = _field(default_factory=dict)
    questions: tuple = ()


def test_round_umbrella_single_size_gets_dome_neck_plus_1000():
    # Решение 09.10.2026: купол = патрубок + 2 × 500 (вылет), как в скилле; было 2 × патрубок.
    position = _Position("1", "Зонт круглый 250", "umbrella_round", 1, size={"d": 250},
                         questions=(_Question("missing_size_d_dome", "Не указан диаметр купола"),))
    fixed = actions.umbrella_with_dome(position)
    assert fixed.size["d_neck"] == 250 and fixed.size["d_dome"] == 1250
    assert fixed.questions == ()
    explicit = _Position("2", "Зонт", "umbrella_round", 1, size={"d_neck": 250, "d_dome": 450})
    assert actions.umbrella_with_dome(explicit).size["d_dome"] == 450
    square = _Position("3", "Зонт 300х300", "umbrella_square", 1, size={"side_neck": 300})
    assert actions.umbrella_with_dome(square) is square


def test_section_heading_detection():
    def heading(text, **fields):
        return actions.is_section_heading(text, _Position("1", text, **fields))

    assert heading("Вытяжка из химлаборатории.")
    assert heading("Спецификация\tНаименование\tКол-во\tЕд.")
    assert heading("Приток П2, административный корпус")
    assert not heading("Переход ф315/ф250")
    assert not heading("Уголок 45х45х4")
    assert not heading("Скотч\t2\tшт")
    assert not heading("Отвод", size={"d": 200})
    assert not heading("Хомут", qty=2)


def test_preview_without_headings_rechecks_completeness():
    from dataclasses import replace

    heading = _Position("1", "Вытяжка из химлаборатории")
    duct = _Position("2", "Воздуховод ⌀200 — 2 шт", "straight_round", 2, 1500, {"d": 200})
    lost = _NS(source_no="3", source_text="Что-то непонятное", position=None, status="problem")
    item = lambda p: _NS(source_no=p.source_no, source_text=p.source_text, position=p, status="ready")  # noqa: E731
    message = "В разборе есть позиции без корректного количества"
    completeness = _NS(source_position_count=3, source_quantity_sum=None, message=message)

    @_dc(frozen=True)
    class Preview:
        positions: tuple
        response: object
        completeness: object
        source_position_count: int = 3
        model_position_count: int = 2
        model_quantity_sum: float = 2
        completeness_message: str | None = message
        warnings: tuple = (message, "другое")

    @_dc(frozen=True)
    class Response:
        positions: tuple

    def compare(source, response):
        count = len(response.positions)
        ok = all(p.qty is not None for p in response.positions)
        return _NS(source_position_count=source.position_count, model_position_count=count,
                   model_quantity_sum=sum(p.qty or 0 for p in response.positions),
                   message=None if ok and count == source.position_count else "расхождение")

    preview = Preview((item(heading), item(duct), lost), Response((heading, duct)), completeness)
    fixed = actions.preview_without_headings(preview, compare)
    assert [p.source_no for p in fixed.positions] == ["2", "3"]  # неразобранная строка не пропала
    assert fixed.response.positions == (duct,)
    # Источник: 3 строки минус заголовок = 2, модель вернула 1 позицию — расхождение видно.
    assert fixed.completeness.source_position_count == 2
    assert fixed.warnings == ("другое", "расхождение")
    assert actions.preview_without_headings(replace(preview, positions=(item(duct),)), compare).positions == (item(duct),)


def test_xlsx_without_index_column_becomes_numbered_lines(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Вытяжка В2, Северный цех. "])
    sheet.append(["Наименование", "Кол-во", "Ед."])
    sheet.append(["Воздуховод из полипропилена ⌀200", 8, "мп"])
    sheet.append(["Зонт круглый 315", 1, "шт"])
    plain = tmp_path / "plain.xlsx"
    book.save(plain)
    assert actions.spec_lines_text(plain).splitlines() == [
        "Вытяжка В2, Северный цех.",
        "1. Воздуховод из полипропилена ⌀200 — 8 мп; кол-во: 8",
        "2. Зонт круглый 315 — 1 шт; кол-во: 1",
    ]
    sheet.insert_cols(1)
    sheet["A2"] = "№"
    numbered = tmp_path / "numbered.xlsx"
    book.save(numbered)
    assert actions.spec_lines_text(numbered) is None  # с «№» позиции делит сам прототип


def test_unit_column_before_quantity_gives_number_then_unit(tmp_path):
    """Прототип узнаёт метраж только как «54.8 м»: при «м. 54.8» прямой участок не режется по 1500,
    а модель ставит 54,8 шт по 54 800 мм. Колонка «Ед.» перед «Кол-вом» — частый вид спецификаций."""

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([None, None, "Система Во 1.1"])
    sheet.append([None, None, "Скруббер STRADA CLEAN C 18-35-750", "ХИМВЕНТ-Н-О-810", None, None, "шт.", 1])
    sheet.append([None, None, "Воздуховод из полипропилена ø150, толщ. 3,0 мм", None, None, None, "м.", 54.8])
    sheet.append([None, None, "Воздуховод из полипропилена ø900, толщ. 5,0 мм", None, None, None, "м.", 0.5])
    path = tmp_path / "unit_first.xlsx"
    book.save(path)
    assert actions.spec_lines_text(path).splitlines() == [
        "1. Скруббер STRADA CLEAN C 18-35-750 — ХИМВЕНТ-Н-О-810 1 шт.; система Во 1.1; кол-во: 1",
        "2. Воздуховод из полипропилена ø150, толщ. 3,0 мм — 54.8 м.; система Во 1.1; кол-во: 54.8",
        "3. Воздуховод из полипропилена ø900, толщ. 5,0 мм — 0.5 м.; система Во 1.1; кол-во: 0.5",
    ]


def test_request_in_copy_of_our_calc_template_keeps_only_name_quantity_unit(tmp_path):
    """Заявка 09.10.2026 пришла в копии нашей расчётки: шапка «Номенклатура … S элемента, м2» с цифрой
    становилась позицией № 1 и сдвигала номера, а нули и #DIV/0! расчётных колонок попадали в каждую строку."""

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Номенклатура ", "Кол-во", "Ед.изм.", "Толщина материала в мм", "S элемента, м2", "S общ ",
                  "Цена  кг  ", "Плотность", "% на раскрой", "К/эф", "Цена за шт без соединения, руб",
                  "ИТОГО без соединения", "СОЕДИНЕНИЕ", "Цена за шт с соединением, руб",
                  "ИТОГО с соединением, руб", "Масса воздуховодов всего, с креплением, кг *"])
    sheet.append(["Дроссель клапан  ПП  400х250 фф", 2, "шт", *[None] * 10, 0, 0, "#DIV/0!"])
    sheet.append(["Шумоглушитель ПП ф500 L 900 фф", 1, "шт", *[None] * 10, 0, 0, "#DIV/0!"])
    path = tmp_path / "client.xlsx"
    book.save(path)
    assert actions.spec_lines_text(path).splitlines() == [
        "1. Дроссель клапан ПП 400х250 фф — 2 шт; кол-во: 2",
        "2. Шумоглушитель ПП ф500 L 900 фф — 1 шт; кол-во: 1",
    ]


def fake_prototype(tmp_path, error: str, raise_line: str | None = None) -> Path:
    """Минимальный duct-calc: pipeline падает с ValueError, как прототип на файле без позиций.

    `raise_line` заменяет эту ошибку другой строкой `raise …` (например, ошибкой клиента модели).
    """

    package = tmp_path / "proto" / "src" / "duct_calc"
    package.mkdir(parents=True)
    (package / "__init__.py").write_text("", encoding="utf-8")
    (package / "model_client.py").write_text(
        "class ModelClientError(RuntimeError): pass\n"
        "class ModelClientConfig:\n"
        "    last = None\n"
        "    def __init__(self, **kw):\n"
        "        ModelClientConfig.last = kw\n"
        "class DeepSeekClient:\n    def __init__(self, config): pass\n",
        encoding="utf-8",
    )
    raise_line = raise_line or f"raise ValueError({error!r})"
    (package / "pipeline.py").write_text(
        "PASSTHROUGH_ITEM_KEYWORDS = ()\nGRID_KEYWORDS = ()\nVALVE_KEYWORDS = ()\n"
        "def build_system_prompt(): return ''\n"
        "def with_umbrella_optional_neck_key(*a, **kw): return a[0] if a else None\n"
        "def preview_specification(*a, **kw): return None\n"
        "def compare_completeness(*a, **kw): return None\n"
        "import httpx\n"
        "from duct_calc.model_client import ModelClientError\n"
        "class PipelineNeedsInput(ValueError): pass\n"
        f"def run_specification_pipeline(**kw):\n    {raise_line}\n",
        encoding="utf-8",
    )
    (tmp_path / "proto" / "data").mkdir()
    return tmp_path / "proto"


def test_duct_calc_explains_spec_without_positions(tmp_path, monkeypatch):
    import sys

    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        monkeypatch.delitem(sys.modules, name)
    proto = fake_prototype(tmp_path, "Нельзя сформировать расчётку без позиций")
    monkeypatch.setattr(sys, "path", list(sys.path))
    spec_path = tmp_path / "spec.docx"
    spec_path.write_bytes(b"")
    spec = Attachment(id=1, name="spec.docx", path=spec_path, size=0)
    result = actions.duct_calc(
        spec=spec, connection_label="Раструб", duct_calc_dir=proto, base_url="u", api_key="sk-test", model="m"
    )
    assert result.text == actions.NO_POSITIONS_TEXT
    from duct_calc.model_client import ModelClientConfig

    # Зависшее соединение видно через 15 с, ответ модели ждём до 2 минут.
    timeout = ModelClientConfig.last["timeout_seconds"]
    assert (timeout.connect, timeout.read) == (15, 120)
    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        del sys.modules[name]


@pytest.mark.parametrize(
    ("raise_line", "no_connection"),
    [
        ("raise ModelClientError('Не удалось подключиться к API модели') from httpx.ConnectTimeout('tls')", True),
        ("raise ModelClientError('Модель не успела ответить.') from httpx.ReadTimeout('read')", True),
        ("raise ModelClientError('API модели вернул HTTP 503')", True),
        ("raise ModelClientError('API модели вернул HTTP 429')", True),
        ("raise ModelClientError('API модели вернул HTTP 401')", False),
        ("raise ModelClientError('Модель вернула пустой ответ')", False),
    ],
)
def test_duct_calc_says_no_connection_only_for_network_failures(tmp_path, monkeypatch, raise_line, no_connection):
    import sys

    from ameri.llm import NO_CONNECTION_MESSAGE

    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        monkeypatch.delitem(sys.modules, name)
    proto = fake_prototype(tmp_path, "", raise_line=raise_line)
    monkeypatch.setattr(sys, "path", list(sys.path))
    spec_path = tmp_path / "spec.docx"
    spec_path.write_bytes(b"")
    spec = Attachment(id=1, name="spec.docx", path=spec_path, size=0)
    try:
        if no_connection:
            result = actions.duct_calc(
                spec=spec, connection_label="Фланец", duct_calc_dir=proto, base_url="u", api_key="sk-test", model="m"
            )
            assert result.text == NO_CONNECTION_MESSAGE
        else:
            # Поломку разбора не выдаём за сбой связи: менеджер увидит настоящую ошибку.
            with pytest.raises(RuntimeError):
                actions.duct_calc(
                    spec=spec, connection_label="Фланец", duct_calc_dir=proto, base_url="u", api_key="sk-test", model="m"
                )
    finally:
        for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
            del sys.modules[name]


def test_chat_model_says_no_connection_and_waits_two_minutes():
    import httpx

    from ameri.llm import NO_CONNECTION_MESSAGE, DeepSeekChat, LlmError, NoConnectionError

    timeouts = []

    def chat(handler):
        def recording(request):
            timeouts.append(request.extensions["timeout"])
            return handler(request)

        client = httpx.Client(transport=httpx.MockTransport(recording))
        return DeepSeekChat(base_url="https://x", api_key="sk-test", model="m", http_client=client)

    def hang(request):
        raise httpx.ConnectTimeout("handshake", request=request)

    with pytest.raises(NoConnectionError, match=NO_CONNECTION_MESSAGE):
        chat(hang).complete([{"role": "user", "content": "?"}])
    assert len(timeouts) == 2
    assert (timeouts[0]["connect"], timeouts[0]["read"]) == (15, 120)

    with pytest.raises(NoConnectionError):
        chat(lambda request: httpx.Response(503)).complete([{"role": "user", "content": "?"}])

    with pytest.raises(LlmError, match="ошибкой 401") as error:
        chat(lambda request: httpx.Response(401)).complete([{"role": "user", "content": "?"}])
    assert not isinstance(error.value, NoConnectionError)


def test_duct_calc_text_wraps_chat_message(monkeypatch, tmp_path):
    captured = {}

    def fake_duct_calc(**kwargs):
        captured["name"] = kwargs["spec"].name
        captured["text"] = kwargs["spec"].path.read_text(encoding="utf-8")
        return actions.ActionResult("Расчётка по файлу «позиции-из-чата.md» готова.")

    monkeypatch.setattr(actions, "duct_calc", fake_duct_calc)
    result = actions.duct_calc_text(
        text="Шибер 160 — 2 шт",
        connection_label="Фланец",
        duct_calc_dir=tmp_path,
        base_url="u",
        api_key="sk-test",
        model="m",
    )
    assert captured == {"name": "позиции-из-чата.md", "text": "Шибер 160 — 2 шт\n"}
    assert result.text.startswith("Расчётка по тексту из чата")


def test_line_without_connection_is_not_overridden_by_chat_choice():
    plain = "Отвод 90 круглый без соединения из полипропилена СМП-ППС-_/_-О90-400"
    assert actions.names_no_connection(plain)
    assert actions.names_no_connection("Переход ПП ф200, соединение отсутствует")
    assert actions.names_no_connection("Воздуховод ф250, соединение не требуется")
    assert actions.names_no_connection("Отвод ф200 без фланцев и без раструбов")
    assert not actions.names_no_connection("Отвод 90 ПП ф400 муфтовое соединение")
    assert not actions.names_no_connection("Воздуховод круглый из полипропилена ф 355 L-2000 мм")


def test_coupling_wording_is_socket_not_a_separate_product():
    assert actions.is_coupling_item("Муфта из полипропилена СМП-ППС-М-355")
    assert actions.is_coupling_item("Муфта соединительная для круглых воздуховодов полипропиленовая 315")
    assert not actions.is_coupling_item("Отвод 90 круглый без соединения из полипропилена СМП-ППС-_/_-О90-400")
    assert not actions.is_coupling_item("Воздуховод круглый из полипропилена ф 355 муфтовое соединение L-2000")
    assert actions.names_socket_connection("Отвод 90 ПП ф400 муфтовое соединение")
    assert not actions.names_socket_connection("Отвод 90 круглый без соединения из полипропилена")


def test_install_counts_coupling_as_one_socket(tmp_path, monkeypatch):
    import sys
    from types import SimpleNamespace

    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        monkeypatch.delitem(sys.modules, name)
    proto = fake_prototype(tmp_path, "не важно")
    monkeypatch.setattr(sys, "path", [str(proto / "src"), *sys.path])
    from duct_calc import pipeline

    seen = {}

    def original_nipple(text):
        return "ниппел" in text.lower()

    def original_global(position, connection_type):
        seen["connection"] = connection_type
        return position

    pipeline._is_nipple_text = original_nipple
    pipeline.with_global_connection = original_global
    actions._install_harness(None)
    assert pipeline._is_nipple_text("Муфта из полипропилена СМП-ППС-М-355")
    assert pipeline._is_nipple_text("ниппель 400")
    assert not pipeline._is_nipple_text("Отвод 90 круглый без соединения")
    pipeline.with_global_connection(SimpleNamespace(source_text="Отвод 90 ПП ф400 муфтовое соединение"), "none")
    assert seen["connection"] == "socket"
    pipeline.with_global_connection(SimpleNamespace(source_text="Отвод 90 круглый без соединения"), "flange")
    assert seen["connection"] == "none"
    pipeline.with_global_connection(SimpleNamespace(source_text="Воздуховод ф400"), "flange")
    assert seen["connection"] == "flange"
    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        del sys.modules[name]


def test_harness_keeps_prototype_meter_cutting_q37(tmp_path, monkeypatch):
    """Q37 (а): метраж режет сам прототип — 1500 мм и остаток; Харнес не подменяет его нарезку."""

    import sys

    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        monkeypatch.delitem(sys.modules, name)
    proto = fake_prototype(tmp_path, "не важно")
    monkeypatch.setattr(sys, "path", [str(proto / "src"), *sys.path])
    from duct_calc import pipeline

    def prototype_cutting(*args, **kwargs):
        return "нарезка прототипа"

    pipeline._meter_segment_rows = prototype_cutting
    actions._install_harness(None)
    assert pipeline._meter_segment_rows is prototype_cutting
    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        del sys.modules[name]


def test_result_notes_umbrella_and_flange20(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([None, 1, "Зонт круглый 315\t1\tшт"])
    sheet.append([None, 2, "Переход ф450/600х300 — труба/фланец20"])
    sheet.append([None, 3, "Отвод ф200 (проблема: Позиция 3: тип не распознан)"])
    sheet.append([None, 4, "Воздуховод ф200 — 2 шт по 1000 мм"])
    path = tmp_path / "r.xlsx"
    book.save(path)
    notes = actions.result_notes(path, "Раструб")
    assert len(notes) == 2
    assert "Проверьте зонты вручную" in notes[0] and "Зонт круглый 315 · 1 · шт" in notes[0]
    assert "шинорейка 20" in notes[1] and "«раструб»" in notes[1] and "фланец20" in notes[1]
    assert "стандартным фланцем" in actions.result_notes(path, "Фланец")[1]
    assert not actions.FLANGE_20.search("Воздуховод ф200")


def test_umbrella_note_depends_on_type(tmp_path):
    """У вытяжного (островного) зонта купола нет — подпись «купол = патрубок + 1000» его вводила в заблуждение."""

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([None, 4, "Зонт вытяжной ПП 1400х1000 Н300 врезка ф315 площадка 700х500 ф — 3 шт; кол-во: 3"])
    sheet.append([None, 5, "Зонт крышный ПП ф500 ф — 1 шт; кол-во: 1"])
    path = tmp_path / "r.xlsx"
    book.save(path)
    island, roof = actions.result_notes(path, "Фланец")[0].splitlines()[1:]
    assert island.startswith("- «Зонт вытяжной ПП 1400х1000 Н300 врезка ф315 площадка 700х500 ф» — островной")
    assert "врезка — L200 и 1 фланец" in island and "купол" not in island
    assert roof == "- «Зонт крышный ПП ф500 ф» — купол взят как патрубок + 1000 мм, если не указан"


def test_manager_message_drops_model_warnings_keeps_code_ones():
    """Свободный текст модели («element_type = null…») — её промежуточный шаг, менеджеру он не нужен;
    сигнал кода о полноте разбора остаётся."""

    @_dc(frozen=True)
    class Preview:
        response: object
        warnings: tuple

    model = "Позиции «Шумоглушитель» — вне каталога типов: element_type = null."
    code = "В разборе есть позиции без корректного количества"
    preview = Preview(_NS(warnings=(model,)), (model, code))
    assert actions.without_model_warnings(preview).warnings == (code,)
    assert actions.manager_warning("Коэффициент L для нестандарт фасон взят из скилла") == (
        "Раскрой нестандартных изделий (зонты, клапаны, шумоглушители) — 1,5 (в шаблоне «от 1,5, уточнять»)."
    )
    assert actions.manager_warning(code) == code


def test_himvent_and_lone_connection_are_read_from_the_line():
    """ХИМВЕНТ на названном элементе считается. Отдельный фланец берёт размер и количество из строки."""

    assert actions.is_himvent_duct("Воздуховод аналог ХИМВЕНТ ø250 L1500")
    assert actions.is_himvent_duct("Отвод 90-ПП-110-3-ММ ХИМВЕНТ")
    assert actions.is_himvent_duct("Переход ХИМВЕНТ-ПП-250/200-3-ММ")
    assert actions.is_himvent_duct("Прямой участок ХИМВЕНТ ф250")
    # Любой элемент вентиляции, который считает расчётка, — не только воздуховод и фасонные изделия.
    for line in (
        "Шибер ХИМВЕНТ ф250",
        "Заслонка шиберная ХИМВЕНТ 400х200",
        "Клапан обратный ХИМВЕНТ ф315",
        "Дроссель-клапан ХИМВЕНТ ф200",
        "Врезка ХИМВЕНТ 400/160",
        "Решётка ХИМВЕНТ ф250",
        "Сетка защитная ХИМВЕНТ ф250",
        "Шумоглушитель ХИМВЕНТ ф400",
        "Ниппель ХИМВЕНТ ф200",
        "Муфта ХИМВЕНТ ф160",
        "Фланец ХИМВЕНТ ф355",
        "Раструб ХИМВЕНТ ф315",
    ):
        assert actions.is_himvent_duct(line), line
    assert not actions.is_himvent_duct("ХИМВЕНТ-Н-О-1130-ПП-БС-У2-30/1000-ПО")
    assert not actions.is_himvent_duct("Вентилятор ХИМВЕНТ-Н-К-О-225")
    assert not actions.is_himvent_duct("Скруббер ХИМВЕНТ")

    assert actions.lone_connection_kind("фланец 355") == "flange"
    assert actions.lone_connection_kind("фланцы 355 — 4 шт") == "flange"
    assert actions.lone_connection_kind("фланец ПП 600х300") == "flange"
    assert actions.lone_connection_kind("фланец 200") == "flange"
    assert actions.lone_connection_kind("раструб 315") == "socket"
    assert actions.lone_connection_kind("фланец 20") is None
    assert actions.lone_connection_kind("ф20") is None
    assert actions.lone_connection_kind("Отвод 90 ф315 фланец") is None

    assert actions.lone_connection_section("фланец 355") == ("round", (355.0,))
    assert actions.lone_connection_section("фланец 355 — 4 шт") == ("round", (355.0,))
    assert actions.lone_connection_section("фланец ПП 600х300") == ("rect", (600.0, 300.0))
    assert actions.stated_piece_quantity(None, "фланец 355") == 1
    assert actions.stated_piece_quantity(None, "фланец 355 — 4 шт") == 4
    assert actions.stated_piece_quantity(4, "фланец 355") == 4
    # Развернутое описание со словом «для» — тоже отдельный фланец (заявка 145845)
    catalog = "Фланец для прямоугольного воздуховода из полипропилена 200х200 мм, толщина для P=3000 Па, цвет серый"
    assert actions.lone_connection_kind(catalog) == "flange"
    assert actions.lone_connection_section(catalog) == ("rect", (200.0, 200.0))
    pnd = (
        "Фланец для прямоугольного воздуховода из полиэтилена низкого давления для наружной "
        "прокладки морозостойкий (до -40°C), стойкий к УФ, 1000х600 мм, толщина для P=3000 Па, цвет черный"
    )
    assert actions.lone_connection_kind(pnd) == "flange"
    assert actions.lone_connection_section(pnd) == ("rect", (1000.0, 600.0))
    assert actions.lone_connection_kind("Раструб для круглого воздуховода 315") == "socket"
    assert actions.lone_connection_kind("Фланец перехода 630х400") is None
    assert actions.lone_connection_kind("Фланцевая заглушка для воздуховода 400") is None


def test_n_header_spec_goes_through_line_converter(tmp_path):
    """Заявка с шапкой «N … Количество потребности» (форма 145845) режется нашим конвертером:
    прототип получает готовые строки с количеством, шапка и нумерация колонок не позиции."""

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Заявка"
    sheet.append(["N", "Наименование ТМЦ", "ГОСТ,ТУ", "Производи-тель", "Ед. изм.", "Количество потребности"])
    sheet.append(["1", "2", "3", "4", "5", "6"])
    sheet.append([1, "Воздуховод из полипропилена прямоугольного сечения фланцевый 400х400 мм", None, None, "шт.", 8])
    sheet.append([2, "Фланец для прямоугольного воздуховода из полипропилена 200х200 мм", None, None, "шт.", 5])
    path = tmp_path / "заявка.xlsx"
    book.save(path)
    lines = actions.spec_lines_text(path)
    assert lines is not None
    assert "1. Воздуховод из полипропилена прямоугольного сечения фланцевый 400х400 мм — 8 шт.; кол-во: 8" in lines
    assert "2. Фланец для прямоугольного воздуховода из полипропилена 200х200 мм — 5 шт.; кол-во: 5" in lines
    assert "2 — 3 4 5 6" not in lines  # строка нумерации колонок — не позиция


def test_prototype_number_header_spec_is_left_to_prototype(tmp_path):
    """Шапка «№» — родная для прототипа: файл уходит ему целиком, конвертер отступает."""

    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["№", "Наименование", "Ед.", "Кол-во"])
    sheet.append([1, "Отвод 90 ф200", "шт", 3])
    path = tmp_path / "spec.xlsx"
    book.save(path)
    assert actions.spec_lines_text(path) is None
    assert actions.stated_piece_quantity("4", "фланец 355") == 4
    assert actions.stated_piece_quantity("4", "фланец 355 — 2 шт") == 4
    assert actions.stated_piece_quantity("нет", "фланец 355 — 4 шт") == 4


def test_install_prices_lone_connection_and_keeps_named_himvent(tmp_path, monkeypatch):
    """Обёртка поверх старого прототипа: воздуховод с ХИМВЕНТ не покупной, фланец 355 — одно соединение."""

    import sys

    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        monkeypatch.delitem(sys.modules, name)
    proto = fake_prototype(tmp_path, "не важно")
    monkeypatch.setattr(sys, "path", [str(proto / "src"), *sys.path])
    from duct_calc import pipeline

    pipeline.PASSTHROUGH_ITEM_KEYWORDS = ("скруббер", "химвент", "вентилятор")

    def old_scrubber(source_text: str) -> bool:
        lower = source_text.lower().replace("ё", "е")
        return any(word in lower for word in pipeline.PASSTHROUGH_ITEM_KEYWORDS)

    def old_rules(position):
        return position

    def old_row(position, coefficients, *args, system=None, connection_type=None, **kwargs):
        return _NS(passthrough=True, connection_type=connection_type, text=position.source_text)

    pipeline._is_scrubber_text = old_scrubber
    pipeline.with_special_catalog_rules = old_rules
    pipeline.position_to_output_row = old_row

    calls = []

    def fake_cost(kind, shape, dimensions):
        calls.append((kind, shape, dimensions))
        return "one"

    monkeypatch.setattr(actions, "connection_unit_cost", fake_cost)
    monkeypatch.setattr(actions, "_output_row", lambda **kwargs: _NS(**kwargs))
    actions._install_harness(None)
    actions._install_harness(None)

    assert not pipeline._is_scrubber_text("Воздуховод аналог ХИМВЕНТ ø250 L1500")
    assert not pipeline._is_scrubber_text("Отвод 90-ПП-110-3-ММ ХИМВЕНТ")
    assert pipeline._is_scrubber_text("ХИМВЕНТ-Н-О-1130-ПП-БС-У2-30/1000-ПО")
    assert pipeline._is_scrubber_text("Вентилятор ХИМВЕНТ-Н-К-О-225")
    assert pipeline._is_scrubber_text("Скруббер ХИМВЕНТ")
    assert pipeline._is_scrubber_text("Воздуховод скруббер ХИМВЕНТ ф250")

    @_dc(frozen=True)
    class Position:
        source_no: str
        source_text: str
        element_type: str | None
        qty: float | None
        length_mm: float | None
        size: dict
        ends: tuple
        questions: tuple

    def position(text, **kwargs):
        data = dict(
            source_no="1",
            source_text=text,
            element_type=None,
            qty=None,
            length_mm=None,
            size={},
            ends=(),
            questions=(("unknown_type", "Тип не распознан"),),
        )
        data.update(kwargs)
        return Position(**data)

    ready = pipeline.with_special_catalog_rules(position("фланец 355"))
    assert ready.questions == ()
    assert ready.qty == 1
    assert ready.element_type is None
    assert ready.size == {"d": 355.0}
    elbow = position("Отвод 90 ф315 фланец", element_type="elbow_round", qty=1, questions=(("keep", "остаётся"),))
    assert pipeline.with_special_catalog_rules(elbow) is elbow

    one = pipeline.position_to_output_row(position("фланец 355"), None, connection_type="socket")
    assert one.purchased is True
    assert one.area_m2 is None
    assert one.quantity == 1
    assert one.connection_cost == "one"
    assert calls == [("flange", "round", (355.0,))]

    from_text_qty = pipeline.position_to_output_row(position("фланец 355", qty="4"), None, connection_type="socket")
    assert from_text_qty.quantity == 4
    assert from_text_qty.purchased is True
    assert from_text_qty.connection_cost == "one"
    assert calls[-1] == ("flange", "round", (355.0,))

    four = pipeline.position_to_output_row(position("фланец 355 — 4 шт"), None, connection_type="none")
    assert four.quantity == 4
    assert calls == [
        ("flange", "round", (355.0,)),
        ("flange", "round", (355.0,)),
        ("flange", "round", (355.0,)),
    ]

    socket = pipeline.position_to_output_row(position("раструб 315"), None, connection_type="flange")
    assert socket.quantity == 1
    assert calls[-1] == ("socket", "round", (315.0,))

    rect = pipeline.position_to_output_row(position("фланец ПП 600х300"), None, connection_type="socket")
    assert rect.quantity == 1
    assert calls[-1] == ("flange", "rect", (600.0, 300.0))

    before = len(calls)
    assert pipeline.position_to_output_row(position("фланец 20"), None, connection_type="flange").passthrough is True
    assert pipeline.position_to_output_row(
        position("Отвод 90 ф315 фланец", element_type="elbow_round", qty=1), None, connection_type="socket"
    ).passthrough is True
    assert len(calls) == before

    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        del sys.modules[name]


def test_split_system_nomenclature_moves_sheet_name_out_of_item():
    """Исправление #25: система — в свой столбец; имя листа и файла туда не пишем."""

    system, name = actions.split_system_nomenclature(
        "Лист1\t\tОтвод 90 гр. ПНД ф280 фф\t\t8"
    )
    assert system is None
    assert name == "Отвод 90 гр. ПНД ф280 фф"

    system, name = actions.split_system_nomenclature(
        "Заявка\t1\tВоздуховод ПНД ф280 L1915 фф\tшт\t1"
    )
    assert system is None
    assert name == "Воздуховод ПНД ф280 L1915 фф"

    system, name = actions.split_system_nomenclature(
        "Лист1\t\tВоздуховод ПНД ф280 L1915 фф\t\t1 — отрезки по 1500 мм"
    )
    assert system is None
    assert name == "Воздуховод ПНД ф280 L1915 фф — отрезки по 1500 мм"

    system, name = actions.split_system_nomenclature(
        "1. Воздуховод ПП ∅400 δ=4 — 14 мп; система В1; кол-во: 14"
    )
    assert system == "В1"
    assert name == "1. Воздуховод ПП ∅400 δ=4 — 14 мп"  # «; кол-во» — служебный хвост, количество в своей колонке

    system, name = actions.split_system_nomenclature(
        "1. Отвод 90 ф315; система цниим.xlsx; кол-во: 8"
    )
    assert system is None
    assert name == "1. Отвод 90 ф315"

    # Отрезок воздуховода: пометка о нарезке остаётся в номенклатуре, но не в системе.
    system, name = actions.split_system_nomenclature(
        "Воздуховод ø150 — 54.8 м.; система Во 1.1 — отрезки по 1500 мм", existing="Во 1.1 — отрезки по 1500 мм"
    )
    assert (system, name) == ("Во 1.1", "Воздуховод ø150 — 54.8 м. — отрезки по 1500 мм")

    system, name = actions.split_system_nomenclature("Отвод 90 ф315 фланец", existing="П1")
    assert system == "П1"
    assert name == "Отвод 90 ф315 фланец"

    system, name = actions.split_system_nomenclature("Отвод 90 ф315 фланец", existing="Лист1")
    assert system is None
    assert name == "Отвод 90 ф315 фланец"

    @_dc(frozen=True)
    class Row:
        system: str | None
        nomenclature: str
        quantity: int

    row = actions.with_system_column(
        Row(system="Лист1", nomenclature="Лист1\t\tШиберная заслонка  ПНД ф280 фф\t\t3", quantity=3)
    )
    assert row.system is None
    assert row.nomenclature == "Шиберная заслонка  ПНД ф280 фф"
    assert row.quantity == 3

    formula, name = actions.split_system_nomenclature("Лист1\t\t=SUM(G3:G70)\t\t")
    assert formula is None
    assert name == "'=SUM(G3:G70)"


def test_item_name_skips_sheet_title_so_gate_and_flange_are_recognized():
    """Заявка «цниим»: «Лист1\\t\\tШиберная заслонка …» ушла красной строкой «тип не распознан» —
    название ищется в начале строки, а там было имя листа."""

    assert actions.item_name("Лист1\t\tШиберная заслонка  ПНД ф280 фф\t\t3") == "Шиберная заслонка  ПНД ф280 фф"
    assert actions.item_name("Спецификация\t1\tВрезка 400/160\tшт\t1") == "Врезка 400/160"
    assert actions.item_name("1\tШибер 160\tшт\t3") == "Шибер 160"  # строка таблицы Word, без листа
    assert actions.item_name("Шибер 160 — 3 шт") == "Шибер 160 — 3 шт"

    assert actions.flanged_special_kind("Лист1\t\tШиберная заслонка  ПНД ф280 фф\t\t3") == "gate"
    assert actions.flanged_special_kind("Обратный клапан ПП ф500 — 1 шт") == "reverse_valve"
    assert actions.flanged_special_kind("7. Клапан обратный d=200 мм — 1 шт") == "reverse_valve"
    assert actions.flanged_special_kind("Лист1\t\tВрезка 400/160\t\t1") == "tap"
    assert actions.flanged_special_kind("Зонт вытяжной ПП 1400х1000 Н300 врезка ф315") is None
    assert actions.flanged_special_kind("Дроссель-клапан ПП ф250") is None
    assert actions.flanged_special_kind("Лист1\t\tОтвод 90 гр. ПНД ф280 фф\t\t8") is None

    assert actions.lone_connection_kind("Лист1\t\tФланец 355\t\t4") == "flange"
    assert actions.lone_connection_section("Лист1\t\tФланец 355\t\t4") == ("round", (355.0,))
    assert actions.is_coupling_item("Лист1\t\tМуфта соединительная 315\t\t2")


def test_install_puts_flanges_on_gate_reverse_valve_and_tap(tmp_path, monkeypatch):
    """Дополнения 5.1–5.3: у шибера 4 фланца, у обратного клапана и врезки по 2 — при любом
    типе соединения файла. Сечение прототип берёт из названия, а не из имени листа."""

    import re
    import sys

    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        monkeypatch.delitem(sys.modules, name)
    proto = fake_prototype(tmp_path, "не важно")
    monkeypatch.setattr(sys, "path", [str(proto / "src"), *sys.path])
    from duct_calc import pipeline

    def core_item_head(source_text):  # как в прототипе: начало строки до первой цифры
        match = re.match(r"^\s*(?:\d+[.)]\s*)?([^\d]*)", source_text)
        return match.group(1).lower() if match else ""

    seen = []

    def core_rules(position):
        seen.append(position.source_text)
        return _Position(position.source_no, position.source_text, size={"d": 280.0})

    def core_row(position, coefficients, *args, system=None, connection_type=None, **kwargs):
        return _NS(connection_type=connection_type, text=position.source_text)

    pipeline._item_head = core_item_head
    pipeline.with_special_catalog_rules = core_rules
    pipeline.position_to_output_row = core_row
    actions._install_harness(None)
    actions._install_harness(None)

    sheet_gate = "Лист1\t\tШиберная заслонка  ПНД ф280 фф\t\t3"
    assert "шибер" in pipeline._item_head(sheet_gate)
    assert "врезк" in pipeline._item_head("Лист1\t\tВрезка 400/160\t\t1")
    assert pipeline._item_head("Шумоглушитель ПП ф500") == "шумоглушитель пп ф"

    ready = pipeline.with_special_catalog_rules(_Position("2", sheet_gate))
    assert seen[-1] == "Шиберная заслонка  ПНД ф280 фф"
    assert ready.source_text == sheet_gate  # номенклатура и система считаются по исходной строке
    assert ready.size == {"d": 280.0}

    for text in (sheet_gate, "Обратный клапан ПП ф500", "Врезка 400/160", "Шибер 160 без соединения"):
        for chosen in ("socket", "none", "flange"):
            row = pipeline.position_to_output_row(_Position("1", text), None, connection_type=chosen)
            assert row.connection_type == "flange", (text, chosen)
    assert pipeline.position_to_output_row(_Position("1", sheet_gate), None).text == "Шиберная заслонка  ПНД ф280 фф"

    for text in ("Отвод 90 ф315", "Дроссель-клапан ПП ф250"):
        assert pipeline.position_to_output_row(_Position("1", text), None, connection_type="socket").connection_type == "socket"
    assert pipeline.position_to_output_row(
        _Position("1", "Отвод 90 круглый без соединения"), None, connection_type="flange"
    ).connection_type == "none"

    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        del sys.modules[name]


def test_only_first_sheet_goes_to_calc_and_preview(tmp_path):
    """В расчёт идёт только первая вкладка Excel: остальные — черновики и справки клиента."""

    import openpyxl

    book = openpyxl.Workbook()
    first = book.active
    first.title = "Заявка"
    first.append(["Наименование", "Кол-во", "Ед."])
    first.append(["Шибер ф250", 2, "шт"])
    extra = book.create_sheet("Черновик")
    extra.append(["Наименование", "Кол-во", "Ед."])
    extra.append(["Воздуховод ф400", 10, "мп"])
    book.active = 1
    path = tmp_path / "two_sheets.xlsx"
    book.save(path)

    assert actions.spec_lines_text(path).splitlines() == ["1. Шибер ф250 — 2 шт; кол-во: 2"]
    rows = actions._xlsx_rows(path)
    assert rows[1][:2] == ["Шибер ф250", "2"]


def test_round_flange_price_from_sheet_by_outer_edge():
    """Круглый фланец не из прайса: лист 4,5 м², число фланцев с листа — по диаметру фланца.

    Цена материала здесь условная (200 ₽/кг), цена компании — только в шаблоне расчётки.
    """

    from decimal import Decimal

    # ⌀370: 7 000 000 / 529² = 25,01 → 25 с листа.
    assert actions.round_flange_price(370, 6, 200) == Decimal("317.95")
    # ⌀420: 7 000 000 / 579² = 20,88 → 21 с листа (с округлением вниз было бы 20 и 397,44).
    assert actions.round_flange_price(420, 6, 200) == Decimal("378.51")
    # ⌀1100, фланец 12 мм: 7 000 000 / 1259² = 4,42 → 4 с листа.
    assert actions.round_flange_price(1100, 12, 200) == Decimal("3974.40")
    # Толщина и цена материала входят в цену листа прямо пропорционально.
    assert actions.round_flange_price(370, 8, 200) == Decimal("423.94")
    assert actions.round_flange_price(370, 6, 100) == Decimal("158.98")
    # С листа не выходит ни одного фланца — новой формулой не считается.
    assert actions.round_flange_price(4000, 12, 200) is None


DUCT_CALC_SRC = Path(__file__).resolve().parents[1] / "duct-calc" / "src"


@pytest.fixture
def real_duct_calc(monkeypatch):
    """Настоящий прототип duct-calc из соседней папки; в публичном репозитории его нет — тест пропускается."""

    import sys

    if not (DUCT_CALC_SRC / "duct_calc" / "flanges.py").is_file():
        pytest.skip("нет прототипа duct-calc")
    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        monkeypatch.delitem(sys.modules, name)
    monkeypatch.setattr(sys, "path", [str(DUCT_CALC_SRC), *sys.path])
    yield
    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        del sys.modules[name]


def test_round_flange_not_in_price_list_uses_sheet_price(real_duct_calc, tmp_path):
    """Только круглый размер не из прайса меняет цену; прайс и прямоугольник — как у прототипа.

    Толщина стенки позиции на фланец не влияет.
    """

    from decimal import Decimal

    import openpyxl
    from duct_calc.flanges import DEFAULT_PRICE_BOOK, ConnectionEnd, flange_cost
    from duct_calc.xlsx_out import MATERIAL_PRICE_CELL, SHEET_NAME

    book = openpyxl.Workbook()
    book.active.title = SHEET_NAME
    book.active[MATERIAL_PRICE_CELL] = 200
    template = tmp_path / "шаблон.xlsx"
    book.save(template)
    assert actions.template_material_price(template) == Decimal("200")
    assert actions.template_material_price(tmp_path / "нет.xlsx") is None

    from duct_calc import flanges

    original = flanges.FlangePriceBook.pair_price
    actions._install_flange_price(Decimal("200"))
    actions._install_flange_price(Decimal("200"))  # повторная установка не оборачивает обёртку

    def price(end, **kwargs):
        return DEFAULT_PRICE_BOOK.pair_price(end, **kwargs)

    d370 = ConnectionEnd.round(370, "flange")
    assert price(d370).price == Decimal("635.90")  # пара = 2 × 317,95
    assert price(d370).source == "calculated"
    assert flange_cost([d370] * 4) == Decimal("1271.80")  # 4 конца — 2 пары
    assert actions.connection_unit_cost("flange", "round", (370,)) == Decimal("317.95")
    # Толщина стенки отвода (3 мм от модели или из строки) не делает фланец тоньше: он 6 мм.
    assert price(d370, explicit_flange_thickness_mm=3) == price(d370)
    assert flange_cost([d370] * 2, explicit_flange_thickness_mm=3) == Decimal("635.90")

    for end in (ConnectionEnd.round(355, "flange"), ConnectionEnd.rect(550, 350, "flange"), ConnectionEnd.rect(600, 300, "flange")):
        assert price(end) == original(DEFAULT_PRICE_BOOK, end), end
        assert price(end, explicit_flange_thickness_mm=3) == original(DEFAULT_PRICE_BOOK, end), end
    assert original(DEFAULT_PRICE_BOOK, ConnectionEnd.rect(550, 350, "flange"), explicit_flange_thickness_mm=3) != (
        price(ConnectionEnd.rect(550, 350, "flange"))
    )  # у прототипа толщина стенки меняла цену прямоугольного фланца не из прайса
    assert price(ConnectionEnd.round(355, "flange")).source == "table"

    # Фланцы 401–449 мм — 6 мм, как до 400; раньше таблица прототипа на них останавливалась с вопросом.
    from duct_calc.thickness import flange_thickness

    assert [flange_thickness(d_mm=d) for d in (400, 420, 449, 450, 800, 1000)] == [6, 6, 6, 8, 10, 12]
    assert flange_thickness(a_mm=420, b_mm=300) == 6
    assert price(ConnectionEnd.round(420, "flange")).price == 2 * actions.round_flange_price(420, 6, 200)

    # Круглый фланец не из прайса — толщиной как в прайсе шаблона: до 800 мм — 6, от 800 — 8,
    # а не 8/10/12 мм по таблице: иначе ⌀600 выходил дороже ⌀630 из прайса.
    assert [actions.round_flange_sheet_thickness(d) for d in (370, 600, 799, 800, 850, 1100)] == [6, 6, 6, 8, 8, 8]
    assert price(ConnectionEnd.round(600, "flange")).price == Decimal("1324.80")  # 12 с листа, 6 мм
    assert price(ConnectionEnd.round(1100, "flange")).price == Decimal("5299.20")  # 4 с листа, 8 мм
    assert price(ConnectionEnd.round(850, "flange")).price == 2 * actions.round_flange_price(850, 8, 200)


def test_himvent_line_is_counted_as_named_element(real_duct_calc):
    """«ХИМВЕНТ» в строке не делает элемент покупным; покупные слова и голый код — по-прежнему без расчёта."""

    actions._install_harness(HARNESS)
    from duct_calc import pipeline

    for line in (
        "Шибер ХИМВЕНТ ф250",
        "Клапан обратный ХИМВЕНТ ф315",
        "Врезка ХИМВЕНТ 400/160",
        "Решётка ХИМВЕНТ ф250",
        "Шумоглушитель ХИМВЕНТ ф400",
        "Ниппель ХИМВЕНТ ф200",
        "Муфта ХИМВЕНТ ф160",
        "Фланец ХИМВЕНТ ф355",
        "Воздуховод аналог ХИМВЕНТ ø250 L1500",
    ):
        assert not pipeline._is_scrubber_text(line), line
    for line in (
        "ХИМВЕНТ-ВО-250",
        "Вентилятор ХИМВЕНТ ВР-80 с фланцем",
        "Клапан КВР ХИМВЕНТ 300",
        "Скруббер ХИМВЕНТ",
    ):
        assert pipeline._is_scrubber_text(line), line


def test_source_line_is_restored_only_when_model_dropped_system():
    """Хвост «; система …» возвращается, только если модель вернула точное начало строки."""

    line = "24. Отвод из полипропилена ø350-ø350, толщ. 4,0 мм — 9 шт.; система Во 1.1; кол-во: 9"
    full = "Отвод из полипропилена ø350-ø350, толщ. 4,0 мм — 9 шт.; система Во 1.1; кол-во: 9"
    assert actions.restored_source_text("Отвод из полипропилена ø350-ø350, толщ. 4,0 мм — 9 шт.", line) == full
    assert actions.restored_source_text("Отвод из полипропилена  ø350-ø350, толщ. 4,0 мм — 9 шт", line) == full
    assert actions.restored_source_text(full, line) is None
    # Обрезано перед количеством — возвращается и количество.
    assert actions.restored_source_text("Отвод из полипропилена ø350-ø350, толщ. 4,0 мм", line) == full
    # Модель изменила текст — её строку не трогаем.
    assert actions.restored_source_text("Отвод ø350, 9 шт.", line) is None
    # В отброшенном нет системы — тоже не трогаем: так таблицы с колонкой «№» остаются как у прототипа.
    assert actions.restored_source_text("1 Воздуховод ø200", "1 Воздуховод ø200 L=1500 мм") is None


def test_system_survives_model_dropping_it_from_numbered_line(real_duct_calc, tmp_path):
    """Модель вернула строку без «; система Во 1.1» — в колонку «Система» система всё равно попадает."""

    actions._install_harness(None)
    actions._install_harness(None)  # повторная установка не оборачивает обёртку
    from duct_calc import pipeline

    lines = [
        "1. Отвод из полипропилена ø350-ø350, толщ. 4,0 мм — 9 шт.; система Во 1.1; кол-во: 9",
        "2. Воздуховод из полипропилена ø150, толщ. 3,0 мм — 54.8 м.; система П 2; кол-во: 54.8",
        "3. Заглушка из полипропилена ø300, толщ. 4,0 мм — 1 шт.; система В 3; кол-во: 1",
    ]
    source = tmp_path / "строки.md"
    source.write_text("\n".join(lines) + "\n", encoding="utf-8")
    model_texts = {
        "1": "Отвод из полипропилена ø350-ø350, толщ. 4,0 мм — 9 шт.",
        "2": "Воздуховод из полипропилена ø150, толщ. 3,0 мм — 54.8 м.; система П 2; кол-во: 54.8",
        "3": "Заглушка из полипропилена ø300, толщ. 4,0 мм",
    }

    class Client:
        def complete_json(self, system_prompt, user_text):
            return {
                "positions": [
                    {"source_no": no, "source_text": text, "element_type": None, "qty": 1, "size": {}, "ends": []}
                    for no, text in model_texts.items()
                ]
            }

    preview = pipeline.preview_specification(source_path=source, connection_type="flange", model_client=Client())
    systems = {
        str(p.source_no): actions.split_system_nomenclature(p.source_text)[0] for p in preview.positions
    }
    assert systems == {"1": "Во 1.1", "2": "П 2", "3": "В 3"}

    # Красная строка «(проблема: …)»: система — в колонке «Система», а не в номенклатуре.
    item = next(p for p in preview.positions if str(p.source_no) == "1")
    row = pipeline._problem_output_row(item, "не указан размер angle")
    assert (row.system, row.nomenclature) == ("Во 1.1", "Отвод из полипропилена ø350-ø350, толщ. 4,0 мм — 9 шт.")
    assert row.quantity == 9


def test_round_tee_branch_is_the_smaller_diameter():
    """Ответвление круглого тройника — меньший диаметр, магистраль — больший; угол и количество не диаметры."""

    assert actions.round_tee_diameters("Тройник из полипропилена ø300-ø300-ø150, толщ. 4,0 мм — 2 шт.") == (300, 150)
    assert actions.round_tee_diameters("Тройник ПП 560--400--560 с отводом 400") == (560, 400)
    assert actions.round_tee_diameters("Тройник ⌀250-⌀160 под углом 45 гр.") == (250, 160)
    assert actions.round_tee_diameters("Тройник ф315/ф250/ф315 - 1шт - раструб/труба/труба") == (315, 250)
    assert actions.round_tee_diameters("Тройник ø500-ø400-ø250") == (500, 250)  # с переходом: магистраль — больший
    assert actions.round_tee_diameters("Тройник ø450-ø450-ø450") == (450, 450)
    assert actions.round_tee_diameters("Тройник ø250; система П 12/14; кол-во: 2") is None
    for text in ("Тройник ф250-90гр", "Тройник ф110", "Тройник ф250 — 12 шт", "Тройник 200х200х200"):
        assert actions.round_tee_diameters(text) is None, text


def test_round_tee_length_uses_smaller_diameter_even_if_model_took_main(real_duct_calc):
    """Модель вернула ответвление = магистраль или пустое — длина всё равно врезка + 300 по меньшему диаметру."""

    actions._install_harness(None)
    actions._install_harness(None)  # повторная установка не оборачивает обёртку
    from duct_calc import pipeline
    from duct_calc.model_schema import parse_model_response

    def tee(no, text, size, length=None):
        raw = {"source_no": no, "source_text": text, "element_type": "tee_round_round", "qty": 1, "size": size,
               "length_mm": length, "ends": [{"role": r, "connection": "flange"} for r in ("main_in", "main_out", "branch")]}
        position = parse_model_response({"positions": [raw]}).positions[0]
        return pipeline.with_standard_tee_total_length(
            pipeline.with_standard_tee_branch_length(pipeline.with_standard_tee_equal_branch(position))
        )

    text = "Тройник из полипропилена ø300-ø300-ø150, толщ. 4,0 мм — 2 шт."
    for size in ({"d": 300, "branch_d": 300}, {"d": 300, "branch_d": None}, {"d": 150, "branch_d": 300}):
        position = tee("1", text, size)
        assert (position.size["d"], position.size["branch_d"], position.length_mm) == (300, 150, 450), size
        assert not position.questions, size
    # Длина из строки или от модели не меняется; без диаметров в строке остаётся ответ модели.
    assert tee("2", text, {"d": 300, "branch_d": 150}, length=900).length_mm == 900
    assert tee("3", "Тройник ф110", {"d": 110, "branch_d": None}).length_mm == 410


def test_long_straight_is_nonstandard_whatever_the_quantity(real_duct_calc):
    """Прямик длиннее 1500 мм — нестандарт: раскрой 1,5, работа 3, и при 1, и при 7 штуках.

    Иначе 7 штук L2000 выходили дешевле за м², чем стандартный кусок 1500, а одна такая штука — вдвое дороже.
    """

    from decimal import Decimal

    actions._install_harness(None)
    actions._install_harness(None)  # повторная установка не оборачивает обёртку
    from duct_calc import pipeline
    from duct_calc.coefficients import CoefficientTable
    from duct_calc.model_schema import parse_model_response

    if not pipeline.DEFAULT_TEMPLATE_PATH.is_file():
        pytest.skip("нет шаблона расчётки")
    coefficients = CoefficientTable.from_workbook(pipeline.DEFAULT_TEMPLATE_PATH)

    def straight(text, size, qty, length):
        element_type = "straight_round" if "d" in size else "straight_rect"
        raw = {"source_no": "1", "source_text": text, "element_type": element_type, "qty": qty, "size": size,
               "length_mm": length, "ends": [{"role": r, "connection": "flange"} for r in ("in", "out")]}
        position = parse_model_response({"positions": [raw]}).positions[0]
        if length is None:
            return position, None
        return position, pipeline.position_to_output_row(position, coefficients, connection_type="flange")

    def lm(row):
        return (Decimal(str(row.l)), Decimal(str(row.m)))

    nonstandard = (Decimal("1.5"), Decimal("3"))
    assert lm(straight("Воздуховод ПНД ф150 L2000 фф — 7 шт", {"d": 150}, 7, 2000)[1]) == nonstandard
    assert lm(straight("Воздуховод ПНД ф150 L2000 фф — 1 шт", {"d": 150}, 1, 2000)[1]) == nonstandard
    assert lm(straight("Воздуховод 500х300 L2000 — 4 шт", {"a": 500, "b": 300}, 4, 2000)[1]) == nonstandard
    # Ровно 1500 и короче — как в шаблоне, правило их не трогает.
    assert lm(straight("Воздуховод ПНД ф150 L1500 фф — 3 шт", {"d": 150}, 3, 1500)[1]) == (Decimal("1.2"), Decimal("2.4"))
    assert lm(straight("Воздуховод ПНД ф150 L1350 фф — 3 шт", {"d": 150}, 3, 1350)[1]) == (Decimal("1.3"), Decimal("2.4"))

    # Погонаж режется по 1500 и остаток — куски не длиннее 1500, коэффициенты обычные.
    position, _ = straight("Воздуховод ПНД ф150 — 3,2 м", {"d": 150}, 1, None)
    rows = pipeline._meter_segment_rows(
        position, coefficients, total_meters=Decimal("3.2"), system=None, connection_type="flange"
    )
    assert [(row.length_mm, lm(row)) for row in rows] == [
        (1500, (Decimal("1.2"), Decimal("2.4"))),
        (200, (Decimal("1.3"), Decimal("2.4"))),
    ]


def _area_of(element_type, length, **size):
    from duct_calc import pipeline
    from duct_calc.model_schema import ModelPosition

    position = ModelPosition(source_no="1", source_text="t", element_type=element_type, qty=1, thickness_mm=None,
                             length_mm=length, size=size, ends=(), questions=())
    return float(pipeline._area(position))


def test_rect_round_transition_area_uses_length_like_excel(real_duct_calc):
    """Переход прямоугольник–круг: длина входит в площадь (K52 эталона = длина / …), а не диаметр.

    Значения — пересчёт листа «расчет площади» эталонной расчётки (блок K51:M56) в LibreOffice.
    """

    actions._install_harness(None)
    actions._install_harness(None)  # повторная установка безопасна
    assert _area_of("transition_rect_round", 600, a=1000, b=600, d_out=800) == 1.97
    assert _area_of("transition_rect_round", 300, a=600, b=400, d_out=500) == 0.70
    assert _area_of("transition_rect_round", 200, a=300, b=200, d_out=250) == 0.27
    # Длина действительно меняет площадь: раньше все три давали одно значение.
    short, long_ = (_area_of("transition_rect_round", n, a=1000, b=600, d_out=800) for n in (200, 1000))
    assert short < 1.97 < long_


def test_offset_rect_2_pairs_a_with_first_shift_like_excel(real_duct_calc):
    """Утка в двух плоскостях: a со сдвигом h1, b со сдвигом h2 (M120 эталона); при a≠b и h1≠h2 это видно."""

    import math

    actions._install_harness(None)
    actions._install_harness(None)

    def excel(a, b, length, h1, h2):
        a, b, length, h1, h2 = (v / 1000 for v in (a, b, length, h1, h2))
        return 2 * (b * math.hypot(length, h2) + a * math.hypot(length, h1) + 0.1 * (a + b))

    for a, b, length, h1, h2 in ((300, 250, 400, 50, 120), (1000, 200, 300, 20, 400)):
        got = _area_of("offset_rect_2", length, a=a, b=b, offset1=h1, offset2=h2)
        assert got == pytest.approx(excel(a, b, length, h1, h2), abs=0.0011)
    swapped = _area_of("offset_rect_2", 300, a=1000, b=200, offset1=400, offset2=20)
    assert swapped != pytest.approx(_area_of("offset_rect_2", 300, a=1000, b=200, offset1=20, offset2=400), abs=0.01)


def test_mass_column_uses_template_formula(real_duct_calc, tmp_path):
    """Колонка S «Масса» — формула шаблона: с «% на раскрой» и весом соединений, как в строке 12 шаблона."""

    from openpyxl import load_workbook

    actions._install_harness(None)
    from duct_calc import pipeline, xlsx_out

    if not pipeline.DEFAULT_TEMPLATE_PATH.is_file():
        pytest.skip("нет шаблона расчётки")
    row = xlsx_out.OutputRow(system=None, nomenclature="Воздуховод", length_mm=1500, quantity=2, thickness_mm=3,
                             area_m2=1.0, l=1.2, m=2.4, connection_cost=500)
    out = xlsx_out.write_workbook(pipeline.DEFAULT_TEMPLATE_PATH, tmp_path / "out.xlsx", [row, row])
    template_sheet = load_workbook(pipeline.DEFAULT_TEMPLATE_PATH)[xlsx_out.SHEET_NAME]
    sheet = load_workbook(out)[xlsx_out.SHEET_NAME]
    for number in (12, 13):
        assert sheet[f"S{number}"].value == template_sheet["S12"].value.replace("12", str(number))


def test_unclear_fragment_is_found_by_marker_patterns():
    patterns = actions.compile_unclear_markers([r"(?<![\wа-яё])ш\s?\d{2}(?!\d)"])
    assert actions.unclear_fragment("Воздуховод 350*350 ш20 2м", patterns) == "ш20"
    assert actions.unclear_fragment("Воздуховод 350*350 Ш 30 2м", patterns) == "Ш 30"
    # Обычные слова и числа не принимаются за маркер: «шт», «ш» без числа, «ш» внутри слова, «ш200».
    for text in ("Воздуховод 350*350 — 2 шт", "Зонт ш", "Воздуховод швеллер20 2м", "Воздуховод ш200 2м", "Воздуховод 350*350 2м"):
        assert actions.unclear_fragment(text, patterns) is None, text


def test_marker_thickness_is_dropped_but_real_thickness_is_kept():
    """«ш20» модель читает как толщину 20 мм — сбрасываем; настоящую толщину («толщ. 4») не трогаем."""

    from dataclasses import dataclass

    @dataclass(frozen=True)
    class Position:
        source_text: str
        thickness_mm: float | None

    patterns = actions.compile_unclear_markers([r"(?<![\wа-яё])ш\s?\d{2}(?!\d)"])
    assert actions.without_marker_thickness(Position("Воздуховод 350*350 ш20 2м", 20), patterns).thickness_mm is None
    assert actions.without_marker_thickness(Position("Воздуховод 350*350 ш20 толщ. 4 мм", 4), patterns).thickness_mm == 4
    assert actions.without_marker_thickness(Position("Воздуховод 350*350 2м", 20), patterns).thickness_mm == 20
    assert actions.without_marker_thickness(Position("Воздуховод 350*350 ш20 2м", None), patterns).thickness_mm is None


def test_red_positions_split_unclear_from_not_calculated(tmp_path):
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append([None, 1, "Воздуховод 350*350 ш20 2м — отрезки по 1500 мм (проблема: УТОЧНИТЬ: что значит «ш20»?)"])
    sheet.append([None, 2, "Переход ПП ф315\t2\tшт (проблема: Позиция 5: не указан размер d_out)"])
    path = tmp_path / "r.xlsx"
    book.save(path)
    assert actions.red_positions(path, kind="unclear") == [
        "«Воздуховод 350*350 ш20 2м — отрезки по 1500 мм» — УТОЧНИТЬ: что значит «ш20»?"
    ]
    assert actions.red_positions(path, kind="problem") == ["«Переход ПП ф315 · 2 · шт» — не указан второй диаметр (выход)"]
    assert len(actions.red_positions(path)) == 2


def test_sh20_line_is_calculated_red_and_marked_clarify(real_duct_calc, tmp_path):
    """«Воздуховод 350*350 ш20 2м»: считается по понятной части (1500 + 500, толщина 4 из таблицы),
    обе строки красные с «УТОЧНИТЬ»; не покупная позиция с 0 ₽ из-за «толщины 20»."""

    from decimal import Decimal

    import openpyxl

    actions._install_harness(HARNESS)
    actions._install_harness(HARNESS)  # повторная установка не наслаивает обёртки
    from duct_calc import pipeline, xlsx_out
    from duct_calc.coefficients import CoefficientTable
    from duct_calc.model_schema import parse_model_response

    if not pipeline.DEFAULT_TEMPLATE_PATH.is_file():
        pytest.skip("нет шаблона расчётки")
    coefficients = CoefficientTable.from_workbook(pipeline.DEFAULT_TEMPLATE_PATH)

    def rows_for(text, thickness):
        raw = {"source_no": "1", "source_text": text, "element_type": "straight_rect", "qty": 2, "size": {"a": 350, "b": 350},
               "thickness_mm": thickness, "length_mm": None, "ends": [{"role": r, "connection": "flange"} for r in ("in", "out")]}
        position = parse_model_response({"positions": [raw]}).positions[0]
        position = actions.without_marker_thickness(position, pipeline._ameri_unclear_patterns)
        return pipeline._output_rows_for_position(
            position, coefficients, meter_source_nos=("1",), system=None, connection_type="flange"
        )

    rows = rows_for("Воздуховод 350*350 ш20 2м", 20)  # модель прочитала «ш20» как толщину 20
    assert [row.length_mm for row in rows] == [1500, 500]
    assert all(not row.purchased and row.thickness_mm == Decimal(4) and row.area_m2 for row in rows)
    assert all(row.problem == "УТОЧНИТЬ: что значит «ш20»?" for row in rows)

    plain = rows_for("Воздуховод 350*350 2м", None)
    assert [row.problem for row in plain] == [None, None]

    out = xlsx_out.write_workbook(pipeline.DEFAULT_TEMPLATE_PATH, tmp_path / "out.xlsx", rows + plain)
    sheet = openpyxl.load_workbook(out)[xlsx_out.SHEET_NAME]
    assert "(проблема: УТОЧНИТЬ: что значит «ш20»?)" in sheet["C12"].value
    assert sheet["C12"].font.color.rgb.endswith("FF0000") and sheet["C13"].font.color.rgb.endswith("FF0000")
    assert "проблема" not in sheet["C14"].value
