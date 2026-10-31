"""Таблица толщин стенок и фланцев (`harness/duct_calc/wall-thickness.xlsx`) — эталон для сверки с кодом."""

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "harness" / "duct_calc" / "wall-thickness.xlsx"
MATERIAL = ROOT / "harness" / "materials" / "05-company-calc.md"
DUCT_CALC_SRC = ROOT / "duct-calc" / "src"

# Строки листа «Лист1»: круглые воздуховоды (по диаметру), прямоугольные (по периметру), фланцы.
ROUND_ROWS, RECT_ROWS, FLANGE_ROWS = range(3, 10), range(14, 21), range(24, 28)


def _table(rows):
    sheet = load_workbook(XLSX)["Лист1"]
    result = []
    for number in rows:
        start, end, thickness = (sheet.cell(number, column).value for column in (1, 2, 3))
        result.append((start, None if end == "и более" else end, thickness))
    return result


def test_xlsx_has_no_personal_metadata():
    props = load_workbook(XLSX).properties
    assert not props.lastModifiedBy


def test_round_table_is_contiguous_and_d250_is_2mm():
    """Диапазоны круглых воздуховодов идут подряд без пропусков и пересечений; ⌀250 — 2 мм (решение заказчика 2026-10-31)."""

    rows = _table(ROUND_ROWS)
    for (_, end, _), (next_start, _, _) in zip(rows, rows[1:]):
        assert next_start == end + 1
    assert rows[-1][1] is None
    by_diameter = lambda d: next(t for s, e, t in rows if s <= d and (e is None or d <= e))
    assert [by_diameter(d) for d in (250, 251, 1000, 1001, 1600, 1601)] == [2, 3, 6, 8, 8, 10]


def test_xlsx_matches_company_calc_material():
    """Числа круглых воздуховодов в xlsx совпадают с таблицей в материале «Как считает расчётка компании»."""

    text = MATERIAL.read_text(encoding="utf-8")
    header = re.search(r"^\| ⌀, мм \|(.+)\|$", text, re.M).group(1).split("|")
    values = re.search(r"^\| Стенка, мм \|(.+)\|$", text, re.M).group(1).split("|")
    expected = []
    for cell, thickness in zip(header, values):
        numbers = [int(n) for n in re.findall(r"\d+", cell)]
        if cell.strip().startswith("до"):
            expected.append((0, numbers[0], int(thickness)))
        elif cell.strip().startswith("больше"):
            expected.append((numbers[0] + 1, None, int(thickness)))
        else:
            expected.append((numbers[0], numbers[1], int(thickness)))
    assert _table(ROUND_ROWS) == expected


@pytest.fixture
def prototype_thickness(monkeypatch):
    import sys

    if not (DUCT_CALC_SRC / "duct_calc" / "thickness.py").is_file():
        pytest.skip("нет прототипа duct-calc")
    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        monkeypatch.delitem(sys.modules, name)
    monkeypatch.setattr(sys, "path", [str(DUCT_CALC_SRC), *sys.path])
    from duct_calc import thickness

    yield thickness
    for name in [n for n in sys.modules if n == "duct_calc" or n.startswith("duct_calc.")]:
        del sys.modules[name]


def test_prototype_thickness_follows_xlsx(prototype_thickness):
    """Код прототипа даёт те же толщины, что xlsx: на границах и внутри диапазонов (не расходятся)."""

    def points(rows):
        for start, end, _ in rows:
            yield from {start, start + 1, end, end + 1} if end else {start, start + 50}

    def expected(rows, value):
        matches = [t for s, e, t in rows if s <= value and (e is None or value <= e)]
        return max(matches) if matches else None

    round_rows, rect_rows = _table(ROUND_ROWS), _table(RECT_ROWS)
    for d in points(round_rows):
        if (want := expected(round_rows, d)) is not None:
            assert int(prototype_thickness.duct_round_thickness(d)) == want, f"⌀{d}"
    for perimeter in points(rect_rows):
        if (want := expected(rect_rows, perimeter)) is not None:
            assert int(prototype_thickness.duct_rect_thickness(perimeter / 4, perimeter / 4)) == want, f"P={perimeter}"
