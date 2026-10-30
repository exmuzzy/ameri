"""Типовые действия, которые Менеджер вызывает в Чате."""

from __future__ import annotations

import re
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET

import httpx
import yaml
from dataclasses import dataclass, field, is_dataclass, replace
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from .llm import DEEPSEEK_TIMEOUT, NO_CONNECTION_MESSAGE, DeepSeekChat, key_problem
from .materials import MAX_SECTIONS, Material, find_sections, sections_context
from .store import Attachment, Message, Usage

TEXT_SUFFIXES = {".md", ".txt", ".csv"}
PARSED_SUFFIXES = {".xlsx", ".docx", ".doc"}
DUCT_SUFFIXES = {".md", ".xlsx", ".docx", ".doc", ".csv", ".odt"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg"}
CONNECTION_TYPES = {"Фланец": "flange", "Раструб": "socket", "Без соединения": "none"}
DUCT_TEMPLATE_NAME = "шаблон расчетки стоимости воздуховодов.xlsx"
MAX_FILE_CHARS = 60_000
# Бюджет истории Чата в символах (~40–50 тыс. токенов): старые сообщения отбрасываются первыми.
MAX_HISTORY_CHARS = 150_000
PREVIEW_MAX_CHARS = 20_000
PREVIEW_MAX_ROWS = 200


@dataclass(frozen=True)
class ActionResult:
    text: str
    files: list[tuple[str, bytes]] = field(default_factory=list)
    usage: Usage | None = None


@dataclass(frozen=True)
class AttachmentPreview:
    """Предпросмотр вложения в Чате: без обращения к DeepSeek или прототипу duct-calc."""

    kind: str  # "image" | "pdf" | "markdown" | "text" | "table" | "unsupported"
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)
    truncated: bool = False


def _odt_node_text(node: ET.Element) -> str:
    parts = [node.text or ""]
    for child in node:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "s":
            parts.append(" " * int(next((v for k, v in child.attrib.items() if k.endswith("}c")), "1")))
        elif tag in ("tab", "line-break"):
            parts.append(" ")
        else:
            parts.append(_odt_node_text(child))
        parts.append(child.tail or "")
    return "".join(parts)


def odt_text(path: Path) -> str:
    """Текст OpenDocument (.odt): абзац — строка, строка таблицы — ячейки через « | »."""

    root = ET.fromstring(zipfile.ZipFile(path).read("content.xml"))
    lines: list[str] = []

    def walk(node: ET.Element) -> None:
        for child in node:
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "table-row":
                cells = [
                    " ".join(_odt_node_text(cell).split())
                    for cell in child
                    if cell.tag.rsplit("}", 1)[-1] == "table-cell"
                ]
                row = " | ".join(cell for cell in cells if cell)
                if row:
                    lines.append(row)
            elif tag in ("p", "h"):
                text = " ".join(_odt_node_text(child).split())
                if text:
                    lines.append(text)
            else:
                walk(child)

    walk(root)
    return "\n".join(lines)


HEADER_WORDS = (
    "наименование", "номенклатура", "кол-во", "количество", "ед.", "ед ", "единица", "позиция", "примечание"
)
INDEX_HEADERS = ("№", "n", "no", "№ п/п", "поз.", "поз")
# Шапки колонки с номерами, которые узнаёт сам прототип duct-calc (document_parser):
# только при них файл уходит ему целиком. Прочие («N», «поз.») прототип не режет
# по номерам сам — их разбирает в строки наш конвертер со значениями формул.
PROTOTYPE_INDEX_HEADERS = ("№", "№ п/п", "номер")
# Расчётные колонки (клиенты присылают заявку в копии нашей расчётки): их нули и #DIV/0! — не текст позиции.
CALC_HEADERS = (
    "s ", "площад", "цена", "плотност", "% на", "раскрой", "к/эф", "коэф", "итого", "соединени", "масса",
    "стоимост", "сумма",
)
EXCEL_ERROR = re.compile(r"#(?:DIV/0!|N/A|VALUE!|REF!|NAME\?|NUM!|NULL!)", re.IGNORECASE)


def _is_header(cells: list[str]) -> bool:
    """Шапка таблицы: несколько ячеек, ни одной с числом, хотя бы одна — «Наименование», «Кол-во», «Ед.»…

    Цифры внутри подписей бывают («S элемента, м2»), поэтому отличаем шапку от позиции по ячейке-числу.
    """

    return (
        len(cells) > 1
        and not any(re.fullmatch(r"\d+(?:[.,]\d+)?", cell) for cell in cells)
        and any(cell.lower().strip().startswith(HEADER_WORDS) or cell.strip().lower() in INDEX_HEADERS for cell in cells)
    )


UNIT_CELL = re.compile(r"(?:м|мп|м\.\s*п|п\.?\s*м|пог\.?\s*м|шт|м2|м²|кв\.?\s*м|компл|к-т|кг)\.?", re.IGNORECASE)


def table_lines(rows: list[list[str]]) -> list[tuple[str, str | None]]:
    """Строки таблицы спецификации — (текст «Наименование — количество ед.», количество).

    Шапка и колонка с номером строки убираются: модель разбора путает их с позициями
    и с количеством. Строка из одной ячейки — заголовок раздела, количество у неё None.
    Колонки шапки с расчётом (площадь, цена, итого) и ошибки Excel в текст не попадают.
    """

    lines: list[tuple[str, str | None]] = []
    index_column = False
    calc_columns: set[int] = set()
    for raw in rows:
        cells = [" ".join(str(cell).split()) for cell in raw]
        if _is_header([cell for cell in cells if cell]):
            index_column = next(cell for cell in cells if cell).lower() in INDEX_HEADERS
            calc_columns = {i for i, cell in enumerate(cells) if cell.lower().startswith(CALC_HEADERS)}
            continue
        cells = [
            cell for i, cell in enumerate(cells) if cell and i not in calc_columns and not EXCEL_ERROR.fullmatch(cell)
        ]
        if not cells:
            continue
        if index_column and len(cells) > 1 and cells[0].rstrip(".").isdigit():
            cells = cells[1:]
        if len(cells) > 1 and all(re.fullmatch(r"\d+(?:[.,]\d+)?", cell) for cell in cells):
            # Строка нумерации колонок («1 | 2 | 3 | 4 | 5 | 6») — не позиция.
            continue
        if len(cells) == 1:
            lines.append((cells[0], None))
            continue
        index = next((i for i, c in enumerate(cells) if i and re.fullmatch(r"\d+(?:[.,]\d+)?", c)), None)
        quantity = cells[index].replace(",", ".") if index else None
        if index and index > 1 and UNIT_CELL.fullmatch(cells[index - 1]):
            # Прототип узнаёт метраж только как «54.8 м», не «м. 54.8».
            cells[index - 1], cells[index] = cells[index], cells[index - 1]
        lines.append((f"{cells[0]} — {' '.join(cells[1:])}", quantity))
    return lines


def _markdown_rows(text: str) -> list[list[str] | str]:
    """Строки Markdown: строка таблицы — список ячеек, остальное — текст без разметки заголовка."""

    items: list[list[str] | str] = []
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("|"):
            cells = [cell.strip() for cell in stripped.strip("|").split("|")]
            if all(set(cell) <= set("-: ") for cell in cells):
                continue  # разделитель |---|---|
            items.append(cells)
        elif stripped:
            items.append(stripped.lstrip("#").strip())
    return items


SYSTEM_HEADING = re.compile(r"^система\s+\S+", re.IGNORECASE)


def _xlsx_items(path: Path) -> list[list[str] | str] | None:
    """Строки Excel без шапки прототипа («№», «№ п/п», «номер»): такие файлы он нарезает
    на позиции сам, по номерам. Остальные — даже с колонкой номеров «N» — читает наш
    конвертер: он берёт вычисленные значения формул Excel, которых прототип не видит.

    """

    import openpyxl

    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - пусть файл читает прототип
        return None
    items: list[list[str] | str] = []
    try:
        # Считается только первая вкладка, как в прототипе: остальные — черновики и справки.
        for sheet in book.worksheets[:1]:
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                filled = [cell for cell in cells if cell]
                if any(cell.lower() in PROTOTYPE_INDEX_HEADERS for cell in filled):
                    return None
                if len(filled) == 1:
                    items.append(filled[0])
                elif filled:
                    items.append(cells)  # с пустыми ячейками: колонки шапки и строк должны совпасть
            items.append("")  # граница листа
    finally:
        book.close()
    return items


def spec_lines_text(path: Path) -> str | None:
    """Спецификация из .md, .odt, .csv или .xlsx без «№» — нумерованными строками для прототипа; None — не наш формат.

    С номерами («1. …») прототип делит текст на позиции сам, без модели; строки без номера
    (заголовок документа, «Система В1») в позиции не попадают. Колонку «СИСТЕМА» прототип
    берёт из строки позиции, поэтому система из заголовка раздела дописывается в каждую строку.
    """

    import csv
    import io

    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t") if text.strip() else csv.excel
        items: list[list[str] | str] = list(csv.reader(io.StringIO(text), dialect))
    elif suffix == ".odt":
        items = [line.split(" | ") if " | " in line else line for line in odt_text(path).splitlines()]
    elif suffix == ".md":
        items = _markdown_rows(path.read_text(encoding="utf-8", errors="replace"))
    elif suffix == ".xlsx":
        items = _xlsx_items(path)
        if items is None:
            return None
    else:
        return None

    rows: list[tuple[str, str | None]] = []
    table: list[list[str]] = []
    for item in [*items, ""]:
        if isinstance(item, list):
            table.append(item)
            continue
        rows += table_lines(table)
        table = []
        if item:
            rows.append((item, None))

    heading: list[str] = []
    lines: list[str] = []
    system = ""
    for text, quantity in rows:
        if quantity is None and " — " not in text:
            # Заголовок — строка вне таблицы или из одной ячейки. Первые — шапка документа
            # до позиций, «Система …» — раздел; остальные в позиции не попадают.
            if SYSTEM_HEADING.match(text):
                system = "система " + text.split(None, 1)[1].rstrip(".:;, ")
            elif not lines:
                heading.append(text)
            continue
        line = f"{len(lines) + 1}. {text}"
        if system:
            line += f"; {system}"  # прототип берёт метку системы до ближайшей «;»
        if quantity is not None:
            line += f"; кол-во: {quantity}"
        lines.append(line)
    return "\n".join(heading + lines)


PROBLEM_FIELDS = {
    "d_out": "второй диаметр (выход)",
    "d_dome": "диаметр купола",
    "length_mm": "длина изделия, мм",
    "branch_d": "диаметр ответвления",
}


def red_positions(xlsx: Path, limit: int = 20) -> list[str]:
    """Позиции, которые расчётка пометила «(проблема: …)», — для ответа в Чате."""

    import re

    import openpyxl

    try:
        sheet = openpyxl.load_workbook(xlsx, read_only=True).worksheets[0]
    except Exception:  # noqa: BLE001 - без списка ответ всё равно полезен
        return []
    found: list[str] = []
    for row in sheet.iter_rows(values_only=True):
        for value in row:
            if not isinstance(value, str) or "(проблема:" not in value:
                continue
            text, _, reason = value.partition("(проблема:")
            text = " ".join(text.replace("\t", " · ").split())
            text = re.sub(r"\s*·\s*Кол-во:.*$", "", text)
            reason = re.sub(r"^\s*Позиция\s*\d*:\s*", "", reason.strip().rstrip(")"))
            for field, title in PROBLEM_FIELDS.items():
                reason = reason.replace(f"размер {field}", title).replace(f"длина {field}", title).replace(field, title)
            found.append(f"«{text[:90]}» — {reason}")
            break
    if len(found) > limit:
        found = found[:limit] + [f"…и ещё {len(found) - limit}"]
    return found


FLANGE_20 = re.compile(r"фланц?\w*\s*-?\s*20\b|\bф\s*20\b|шинорейк\w*\s*20", re.IGNORECASE)


def _result_names(xlsx: Path) -> list[str]:
    """Наименования строк расчётки (колонка C под шапкой шаблона), без красных."""

    import openpyxl

    try:
        sheet = openpyxl.load_workbook(xlsx, read_only=True).worksheets[0]
    except Exception:  # noqa: BLE001 - заметки необязательны
        return []
    names = []
    for row in sheet.iter_rows(values_only=True):
        if len(row) > 2 and isinstance(row[1], int) and isinstance(row[2], str) and "(проблема:" not in row[2]:
            names.append(" ".join(row[2].replace("\t", " · ").split()))
    return names


def _umbrella_note(name: str) -> str:
    """Как посчитан зонт: вытяжной — островной по размерам из строки, остальные — с куполом по Q38."""

    lower = name.lower()
    short = name.split(" — ")[0][:90]
    if "вытяжн" in lower or "остров" in lower:
        note = "островной: низ, площадка (верх) и высота из строки"
        if "врезк" in lower:
            note += ", врезка — L200 и 1 фланец"
        return f"- «{short}» — {note}"
    return f"- «{short}» — купол взят как патрубок + 1000 мм, если не указан"


def result_notes(xlsx: Path, connection_label: str) -> list[str]:
    """Что менеджер проверяет сам после расчётки: зонты (Q38) и «фланец 20» (Q40)."""

    names = _result_names(xlsx)
    notes = []
    umbrellas = [n for n in names if "зонт" in n.lower()]
    if umbrellas:
        notes.append("Проверьте зонты вручную:\n" + "\n".join(_umbrella_note(n) for n in umbrellas))
    flanges = [n for n in names if FLANGE_20.search(n)]
    if flanges:
        where = (
            "посчитано со стандартным фланцем компании для этого сечения"
            if connection_label == "Фланец"
            else f"посчитано с выбранным соединением «{connection_label.lower()}»"
        )
        notes.append(
            f"Уточните у клиента «фланец 20» — шинорейка 20, толщина фланца 20 или полка 20? Пока {where}:\n"
            + "\n".join(f"- «{n[:90]}»" for n in flanges)
        )
    return notes


def duct_calc_available(duct_calc_dir: Path | None) -> bool:
    return bool(
        duct_calc_dir
        and (duct_calc_dir / "src" / "duct_calc").is_dir()
        and (duct_calc_dir / "data" / DUCT_TEMPLATE_NAME).is_file()
    )


def _import_duct_calc(duct_calc_dir: Path) -> None:
    src = str(duct_calc_dir / "src")
    if src not in sys.path:
        sys.path.insert(0, src)


def attachment_text(attachment: Attachment, duct_calc_dir: Path | None) -> str:
    """Текст вложения для контекста модели; для неизвестных форматов — только имя."""

    suffix = attachment.path.suffix.lower()
    text = ""
    if suffix in TEXT_SUFFIXES:
        text = attachment.path.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".odt":
        try:
            text = odt_text(attachment.path)
        except (zipfile.BadZipFile, KeyError, ET.ParseError):
            text = ""
    elif suffix in PARSED_SUFFIXES and duct_calc_available(duct_calc_dir):
        _import_duct_calc(duct_calc_dir)
        from duct_calc.document_parser import parse_document

        try:
            text = parse_document(attachment.path).text
        except Exception:  # noqa: BLE001 - файл просто не попадёт в контекст
            text = ""
    if not text:
        return f"[Файл {attachment.name}: содержимое недоступно ассистенту]"
    return f"[Файл {attachment.name}]\n{text[:MAX_FILE_CHARS]}"


def _docx_text(path: Path) -> str:
    from docx import Document

    return "\n".join(p.text for p in Document(path).paragraphs)


def _xlsx_rows(path: Path) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(max_row=PREVIEW_MAX_ROWS + 1, values_only=True)
        ]
    finally:
        workbook.close()


def attachment_preview(attachment: Attachment) -> AttachmentPreview:
    """Предпросмотр вложения в Чате, независимо от прототипа duct-calc.

    Картинки и PDF показываются как есть, .md — форматированным Markdown,
    остальные текстовые форматы и .odt — простым текстом, .xlsx — таблицей;
    для .doc (старый бинарный формат) и прочего показывается только сообщение
    о недоступности предпросмотра.
    """

    suffix = attachment.path.suffix.lower()
    try:
        if suffix in IMAGE_SUFFIXES:
            return AttachmentPreview(kind="image")
        if suffix == ".pdf":
            return AttachmentPreview(kind="pdf")
        if suffix == ".md":
            text = attachment.path.read_text(encoding="utf-8", errors="replace")
            return AttachmentPreview(kind="markdown", text=text[:PREVIEW_MAX_CHARS], truncated=len(text) > PREVIEW_MAX_CHARS)
        if suffix in TEXT_SUFFIXES:
            text = attachment.path.read_text(encoding="utf-8", errors="replace")
            return AttachmentPreview(kind="text", text=text[:PREVIEW_MAX_CHARS], truncated=len(text) > PREVIEW_MAX_CHARS)
        if suffix == ".odt":
            text = odt_text(attachment.path)
            return AttachmentPreview(kind="text", text=text[:PREVIEW_MAX_CHARS], truncated=len(text) > PREVIEW_MAX_CHARS)
        if suffix == ".docx":
            text = _docx_text(attachment.path)
            return AttachmentPreview(kind="text", text=text[:PREVIEW_MAX_CHARS], truncated=len(text) > PREVIEW_MAX_CHARS)
        if suffix == ".xlsx":
            rows = _xlsx_rows(attachment.path)
            return AttachmentPreview(kind="table", rows=rows[:PREVIEW_MAX_ROWS], truncated=len(rows) > PREVIEW_MAX_ROWS)
    except Exception:  # noqa: BLE001 - предпросмотр не должен ронять страницу
        return AttachmentPreview(kind="unsupported")
    return AttachmentPreview(kind="unsupported")


def ask_assistant(
    llm: DeepSeekChat,
    system_prompt: str,
    history: list[Message],
    duct_calc_dir: Path | None,
    names: dict[str, str],
    materials: list[Material] | None = None,
) -> ActionResult:
    """Свободный вопрос: переписка Чата с текстом вложений в пределах бюджета.

    Системный промпт (Харнес) идёт первым и одинаков для всех Чатов — это
    кэшируемый префикс DeepSeek; переписка добавляется после него. Разделы Материалов
    отрасли, подходящие к двум последним вопросам (и к файлам последнего сообщения),
    дописываются в конец последнего сообщения пользователя: так они не сдвигают
    кэшируемое начало запроса.
    """

    turns: list[dict[str, str]] = []
    files_text = ""  # вложения последнего сообщения пользователя — для подбора Материалов
    for message in history:
        content = message.content
        if message.role == "user":
            content = f"{names.get(message.author, message.author)}: {content}"
        elif message.author != "assistant":
            content = f"[Ответ исправлен руководителем {names.get(message.author, message.author)}]\n{content}"
        attached = [attachment_text(attachment, duct_calc_dir) for attachment in message.attachments]
        for text in attached:
            content += "\n\n" + text
        if message.role == "user":
            files_text = "\n".join(attached)
        turns.append({"role": message.role, "content": content})

    kept: list[dict[str, str]] = []
    budget = MAX_HISTORY_CHARS
    for turn in reversed(turns):
        budget -= len(turn["content"])
        if budget < 0 and kept:
            break
        kept.append(turn)
    kept.reverse()
    if len(kept) < len(turns):
        kept.insert(0, {"role": "user", "content": f"[Ранние сообщения чата ({len(turns) - len(kept)}) опущены из-за длины.]"})

    questions = [m.content for m in history if m.role == "user"][-2:]
    sections = find_sections(materials or [], "\n".join(questions))
    if files_text and len(sections) < MAX_SECTIONS:
        # Сначала — разделы по самому вопросу, остальные места — по тексту приложенного файла.
        extra = [s for s in find_sections(materials or [], files_text) if s not in sections]
        sections += extra[: MAX_SECTIONS - len(sections)]
    if sections:
        reference = sections_context(sections)
        if kept and kept[-1]["role"] == "user":
            kept[-1] = {"role": "user", "content": kept[-1]["content"] + "\n\n" + reference}
        else:
            kept.append({"role": "user", "content": reference})

    text, usage = llm.complete([{"role": "system", "content": system_prompt}, *kept])
    return ActionResult(text=text, usage=usage)


def duct_parse_rules(harness_dir: Path | None) -> str:
    """Правила разбора из Харнеса, которые дописываются к промпту прототипа."""

    if harness_dir is None:
        return ""
    path = harness_dir / "duct_calc" / "parse_rules.md"
    return path.read_text(encoding="utf-8").strip() if path.is_file() else ""


def duct_params(harness_dir: Path | None) -> dict:
    if harness_dir is None:
        return {}
    path = harness_dir / "duct_calc" / "params.yaml"
    return (yaml.safe_load(path.read_text(encoding="utf-8")) or {}) if path.is_file() else {}


_KEYWORD_LISTS = {
    "passthrough_keywords": "PASSTHROUGH_ITEM_KEYWORDS",
    "grid_keywords": "GRID_KEYWORDS",
    "valve_keywords": "VALVE_KEYWORDS",
}

# Название изделия — текст до первого размера. «Муфта … М-355» — само изделие.
# «муфтовое» в хвосте отвода или воздуховода — тип соединения, не отдельная муфта.
_ITEM_HEAD_RE = re.compile(r"^\s*(?:\d+[.)]\s*)?([^\d]*)")
_COUPLING_START_RE = re.compile(r"муфт(?!ов)", re.IGNORECASE)
_COUPLING_RE = re.compile(r"(^|[\s,(])муфт(?!ов)", re.IGNORECASE)
_OTHER_PRODUCT_RE = re.compile(
    r"воздуховод|прямик|отвод|тройник|переход|зонт|колено|заглуш|шибер|клапан|реш[её]т|врезк|утк|дефлектор|шумоглуш",
    re.IGNORECASE,
)


def item_name(source_text: str) -> str:
    """Название изделия из строки таблицы: без имени листа Excel, номера и количества.

    Прототип пишет строку .xlsx как «Лист1\\t№\\tНазвание\\t…\\tКол-во», строку таблицы
    Word — «№\\tНазвание\\t…». Название ищется в начале строки, до первой цифры, поэтому
    с именем листа впереди шибер или фланец не узнаётся. Строка без табуляции не меняется.
    """

    text = source_text or ""
    if "\t" not in text:
        return text
    fields = [field.strip() for field in text.split("\t")]
    names = [field for field in fields if _is_nomenclature_field(field)]
    if not names:
        return text
    sheet_first = len(names) > 1 and names[0] == fields[0] and (
        _DEFAULT_SHEET_RE.fullmatch(fields[0]) or not fields[1] or fields[1].isdigit()
    )
    return names[1] if sheet_first else names[0]


def item_head(source_text: str) -> str:
    """Начало названия изделия до первого размера, в нижнем регистре."""

    match = _ITEM_HEAD_RE.match(_plain_item_text(item_name(source_text)))
    return (match.group(1) if match else "").lower()


def flanged_special_kind(source_text: str) -> str | None:
    """Шибер, обратный клапан и врезка — фланцы у них при любом типе соединения файла.

    Дополнения заказчика к расчёту, пункты 5.1–5.3: шибер — 4 фланца, обратный клапан
    и врезка — по 2. Узнаём по названию, как прототип: «Зонт … врезка ф315» — не врезка.
    """

    head = item_head(source_text)
    if "шибер" in head:
        return "gate"
    if "обратн" in head and "клапан" in head:
        return "reverse_valve"
    if "врезк" in head:
        return "tap"
    return None


def _with_source_text(position, text: str):
    if not is_dataclass(position) or getattr(position, "source_text", None) == text:
        return position
    return replace(position, source_text=text)


def is_coupling_item(source_text: str) -> bool:
    """Строка сама называется муфтой: одно раструбное соединение, как ниппель.

    «Муфта … для воздуховодов» остаётся муфтой. «Воздуховод … муфтовое» — нет:
    там муфта только тип соединения другого изделия.
    """

    head = _ITEM_HEAD_RE.match(item_name(source_text))
    text = (head.group(1) if head else "").replace("ё", "е").strip()
    if _COUPLING_START_RE.match(text):
        return True
    if _OTHER_PRODUCT_RE.search(text):
        return False
    return _COUPLING_RE.search(text) is not None


def names_socket_connection(source_text: str) -> bool:
    """«Муфта» и «муфтовое соединение» — то же, что раструб."""

    return "муфт" in (source_text or "").lower().replace("ё", "е")


# «Без соединения» и близкие формулировки на самой позиции. Код «_/_» — то же у СМП.
_NO_CONNECTION_RE = re.compile(
    r"без\s+соединен"
    r"|без\s+соединител"
    r"|соединен\w*\s+отсутств"
    r"|соединен\w*\s+не\s+(?:требуется|предусмотр)"
    r"|без\s+фланц\w*(?:\s+и\s+|\s*/\s*)без\s+раструб"
    r"|без\s+раструб\w*(?:\s+и\s+|\s*/\s*)без\s+фланц"
    r"|_/_",
    re.IGNORECASE,
)


def names_no_connection(source_text: str) -> bool:
    """Позиция сама помечена как изделие без фланца и без раструба."""

    return _NO_CONNECTION_RE.search((source_text or "").replace("ё", "е")) is not None


def _plain_item_text(source_text: str) -> str:
    return (source_text or "").replace("ё", "е").replace("\u00a0", " ").replace("\u202f", " ")


# «ХИМВЕНТ» само по себе — без формулы. Рядом с названием элемента строка считается.
_HIMVENT_DUCT_RE = re.compile(
    r"воздуховод|прям\w*\s+участ|прямик|отвод|колен|переход|тройник|заглуш|утк|зонт|дефлектор"
    r"|шибер|заслонк|клапан|дроссел|врезк|решетк|сетк|шумоглуш|глушител|ниппел|муфт|фланец|фланц|раструб",
    re.IGNORECASE,
)
_LONE_CONNECTION_FILLER = frozenset(
    {
        "пп",
        "полипропилен",
        "полипропиленовый",
        "полипропиленовая",
        "полипропиленовое",
        "круглый",
        "круглая",
        "круглое",
        "прямоугольный",
        "прямоугольная",
        "прямоугольное",
        "плоский",
        "плоская",
        "плоское",
        "из",
        "d",
        "ф",
        "д",
    }
)
_FLANGE_WORDS = frozenset({"фланец", "фланцы"})
_SOCKET_WORDS = frozenset({"раструб", "раструбы"})
# «фланец 20» и «ф20» — не диаметр 20 мм. «фланец 200» сюда не попадает.
_FLANGE_20_RE = re.compile(
    r"(?:фланец|фланц\w*)(?:\s*|-)?20(?!\d)|\bф\s*20(?!\d)",
    re.IGNORECASE,
)
_PIECE_QTY_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*шт\.?", re.IGNORECASE)
_LONE_RECT_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*[xх×]\s*(\d+(?:[.,]\d+)?)", re.IGNORECASE)
_LONE_MARKED_ROUND_RE = re.compile(r"(?:д|ø|⌀|∅)\s*(\d+(?:[.,]\d+)?)", re.IGNORECASE)
_SYSTEM_MARK_RE = re.compile(r"\bсистема\s+([^,\n;]+)", re.IGNORECASE)
# Прототип duct-calc для .xlsx пишет «ИмяЛиста\\tстрока». Заметка нарезки метража — суффикс после « — ».
_SYSTEM_TAIL_RE = re.compile(r";\s*система\s+[^;\n]+", re.IGNORECASE)
_SEGMENT_NOTE_RE = re.compile(r"\s+—\s+(?:отрезки по |остаток ).*$", re.IGNORECASE)


def is_himvent_duct(source_text: str) -> bool:
    """«ХИМВЕНТ» в строке, где назван считаемый элемент воздуховода.

    Голый код «ХИМВЕНТ-…» и вентилятор с этой маркой сюда не входят.
    """

    text = _plain_item_text(source_text).lower()
    return "химвент" in text and _HIMVENT_DUCT_RE.search(text) is not None


def lone_connection_kind(source_text: str) -> str | None:
    """Строка сама называется фланцем или раструбом и задаёт размер.

    «Отвод … фланец» — изделие с соединением на хвосте. «фланец 20» — не диаметр.
    Развернутое описание со словом «для» («Фланец для прямоугольного воздуховода
    из полипропилена 200х200 мм, толщина …») — тоже отдельный фланец.
    """

    text = _plain_item_text(item_name(source_text))
    if _FLANGE_20_RE.search(text):
        return None
    head_match = _ITEM_HEAD_RE.match(text)
    head = (head_match.group(1) if head_match else "").lower().replace("ё", "е")
    words = re.findall(r"[a-zа-я]+", head)
    core = [word for word in words if word not in _LONE_CONNECTION_FILLER]
    if len(core) == 1 and core[0] in _FLANGE_WORDS:
        return "flange"
    if len(core) == 1 and core[0] in _SOCKET_WORDS:
        return "socket"
    if words and "для" in words:
        if words[0] in _FLANGE_WORDS:
            return "flange"
        if words[0] in _SOCKET_WORDS:
            return "socket"
    return None


def stated_piece_quantity(qty: object, source_text: str) -> float:
    """Количество штук из поля позиции, иначе из «N шт» в строке, иначе 1."""

    if isinstance(qty, str):
        text = qty.strip().replace(",", ".")
        if re.fullmatch(r"\d+(?:\.\d+)?", text):
            qty = float(text)
    if isinstance(qty, (int, float)) and not isinstance(qty, bool) and qty >= 1:
        return float(qty)
    piece = _PIECE_QTY_RE.search(_plain_item_text(source_text))
    if piece is not None:
        return float(piece.group(1).replace(",", "."))
    return 1.0


def lone_connection_section(source_text: str) -> tuple[str, tuple[float, ...]] | None:
    """Сечение отдельного фланца или раструба: прямоугольник либо одно число-диаметр.

    Число перед «шт» — количество, не диаметр.
    """

    text = _plain_item_text(item_name(source_text))
    rect = _LONE_RECT_RE.search(text)
    if rect is not None:
        return (
            "rect",
            (float(rect.group(1).replace(",", ".")), float(rect.group(2).replace(",", "."))),
        )
    marked = _LONE_MARKED_ROUND_RE.search(text)
    if marked is not None:
        return ("round", (float(marked.group(1).replace(",", ".")),))
    piece = _PIECE_QTY_RE.search(text)
    qty_span = piece.span(1) if piece is not None else None
    numbers: list[float] = []
    for match in re.finditer(r"\d+(?:[.,]\d+)?", text):
        if qty_span is not None and match.span() == qty_span:
            continue
        numbers.append(float(match.group(0).replace(",", ".")))
    if len(numbers) == 1:
        return ("round", (numbers[0],))
    return None


def _system_label_from_text(text: str) -> str | None:
    match = _SYSTEM_MARK_RE.search(text or "")
    if match is None:
        return None
    label = re.sub(r"\s+", " ", match.group(1)).strip(" .,:;")
    return label or None


def _is_qty_or_unit_field(field: str) -> bool:
    text = field.strip()
    return bool(re.fullmatch(r"\d+(?:[.,]\d+)?", text) or UNIT_CELL.fullmatch(text))


def _is_header_field(field: str) -> bool:
    low = field.lower().strip()
    return low in INDEX_HEADERS or any(low.startswith(word) for word in HEADER_WORDS + CALC_HEADERS)


def _is_nomenclature_field(field: str) -> bool:
    """Ячейка названия изделия: есть буквы, это не шапка, не «8» и не «шт»."""

    text = field.strip()
    if not text or _is_qty_or_unit_field(text) or _is_header_field(text) or EXCEL_ERROR.fullmatch(text):
        return False
    return bool(re.search(r"[а-яa-z]", text, re.IGNORECASE))


_SOURCE_FILENAME_RE = re.compile(
    r"(?i)(?:^|[/\\])[^/\\]+\.(?:xlsx|xls|xlsm|csv|odt|docx|doc|md)$"
)
_DEFAULT_SHEET_RE = re.compile(r"(?i)^(?:лист|sheet)\s*\d+$")


def _is_sheet_or_source_filename(label: str | None) -> bool:
    """Имя листа Excel и имя исходного файла — не обозначение системы."""

    text = (label or "").strip()
    if not text:
        return False
    return bool(_DEFAULT_SHEET_RE.fullmatch(text) or _SOURCE_FILENAME_RE.search(text))


def _allowed_system(label: str | None) -> str | None:
    text = (label or "").strip()
    if not text or _is_sheet_or_source_filename(text):
        return None
    return text


def split_system_nomenclature(text: str, existing: str | None = None) -> tuple[str | None, str]:
    """Настоящая система — в столбец «Система»; из номенклатуры её убрать.

    Имя листа Excel и имя исходного файла системой не бывают: из номенклатуры
    их тоже убираем, в колонку «Система» не пишем. В выгрузке прототипа имя
    листа — первое поле до табуляции. Метка «; система В1» из заголовка раздела
    остаётся обозначением системы.
    """

    raw = _plain_item_text(text).strip()
    if not raw:
        return (_allowed_system(existing), raw)
    system = _allowed_system(existing)
    note_match = _SEGMENT_NOTE_RE.search(raw)
    note = note_match.group(0) if note_match else ""
    body = raw[: note_match.start()].rstrip() if note_match else raw
    if "\t" in body:
        fields = body.split("\t")
        prefix = fields[0].strip()
        name = next((field.strip() for field in fields[1:] if _is_nomenclature_field(field)), "")
        if name and prefix and prefix.casefold() != name.casefold():
            # Первое поле выгрузки xlsx — имя листа, не система.
            if system and system.casefold() == prefix.casefold():
                system = None
        elif prefix and not name:
            name = prefix
        name = name or body
    else:
        if not system:
            system = _allowed_system(_system_label_from_text(body))
        name = _SYSTEM_TAIL_RE.sub("", body).strip(" ;") or body
    if note and note.strip() not in name:
        name = name.rstrip() + note
    return _allowed_system(system), _excel_plain_text(name)


def _excel_plain_text(name: str) -> str:
    """Текст, который Excel не примет за формулу: иначе книга не открывается."""

    if name[:1] in "=+-@":
        return "'" + name
    return name


def with_system_column(row):
    """Систему из номенклатуры — в поле system; имя листа и файла туда не ставить."""

    nomenclature = getattr(row, "nomenclature", None)
    if not isinstance(nomenclature, str):
        return row
    system, name = split_system_nomenclature(nomenclature, getattr(row, "system", None))
    if system == getattr(row, "system", None) and name == nomenclature:
        return row
    if is_dataclass(row) and not isinstance(row, type):
        return replace(row, system=system, nomenclature=name)
    return row


def connection_unit_cost(kind: str, shape: str, dimensions: tuple[float, ...]):
    """Цена одного соединения. Считает прототип, свою цену сюда не пишем."""

    from duct_calc.flanges import ConnectionEnd, connection_cost

    if shape == "round":
        end = ConnectionEnd.round(dimensions[0], kind)
    else:
        end = ConnectionEnd.rect(dimensions[0], dimensions[1], kind)
    return connection_cost((end,))


def _output_row(**kwargs):
    from duct_calc.xlsx_out import OutputRow

    return OutputRow(**kwargs)


def lone_connection_output_row(position, kind: str, section: tuple[str, tuple[float, ...]], *, system: str | None):
    """Строка расчётки: площади тела нет, в соединении — цена одного фланца или раструба."""

    shape, dimensions = section
    text = getattr(position, "source_text", "") or ""
    row_system, name = split_system_nomenclature(
        text, system if system is not None else _system_label_from_text(text)
    )
    source_no = getattr(position, "source_no", "")
    return _output_row(
        system=row_system,
        nomenclature=name.strip() or f"Позиция {source_no}",
        length_mm=None,
        quantity=stated_piece_quantity(getattr(position, "qty", None), text),
        thickness_mm=None,
        area_m2=None,
        l=None,
        m=None,
        connection_cost=connection_unit_cost(kind, shape, dimensions),
        # Флаг гасит тело изделия (толщина, площадь, L и M). Цена соединения
        # пишется в колонку P и входит в итог строки (N+P).
        purchased=True,
    )


def as_lone_connection(position):
    """Снять вопросы о типе и записать количество, чтобы файл с одним фланцем собрался."""

    text = getattr(position, "source_text", "") or ""
    if lone_connection_kind(text) is None:
        return position
    section = lone_connection_section(text)
    if section is None or not is_dataclass(position):
        return position
    shape, dimensions = section
    size = {"d": dimensions[0]} if shape == "round" else {"a": dimensions[0], "b": dimensions[1]}
    return replace(
        position,
        element_type=None,
        size=size,
        length_mm=None,
        ends=(),
        questions=(),
        qty=stated_piece_quantity(getattr(position, "qty", None), text),
    )


def _install_coupling_as_socket(pipeline) -> None:
    """Муфта — один раструб. «Без соединения» на строке гасит и фланец, и раструб.

    Прототип на сервере не меняется: слова ищутся здесь. Выбор в окне чата
    не перекрывает пометку на конкретной позиции.
    """

    if hasattr(pipeline, "_is_nipple_text"):
        original = getattr(pipeline, "_ameri_original_is_nipple_text", None) or pipeline._is_nipple_text
        pipeline._ameri_original_is_nipple_text = original

        def _is_nipple_text(source_text: str) -> bool:
            return original(source_text) or is_coupling_item(source_text)

        pipeline._is_nipple_text = _is_nipple_text

    if hasattr(pipeline, "with_global_connection"):
        original_global = (
            getattr(pipeline, "_ameri_original_global_connection", None) or pipeline.with_global_connection
        )
        pipeline._ameri_original_global_connection = original_global

        def with_global_connection(position, connection_type):
            text = getattr(position, "source_text", "") or ""
            if names_no_connection(text):
                connection_type = "none"
            elif names_socket_connection(text):
                connection_type = "socket"
            return original_global(position, connection_type)

        pipeline.with_global_connection = with_global_connection

    if hasattr(pipeline, "position_to_output_row"):
        original_row = (
            getattr(pipeline, "_ameri_original_position_to_output_row", None) or pipeline.position_to_output_row
        )
        pipeline._ameri_original_position_to_output_row = original_row

        def position_to_output_row(position, coefficients, *args, system=None, connection_type=None, **kwargs):
            text = getattr(position, "source_text", "") or ""
            kind = lone_connection_kind(text)
            if kind is not None:
                section = lone_connection_section(text)
                if section is not None:
                    # Тип соединения берётся из строки. Выбор на весь файл эту строку не перебивает.
                    return lone_connection_output_row(position, kind, section, system=system)
            if flanged_special_kind(text) is not None:
                # Фланцы шибера, обратного клапана и врезки — часть изделия: ни выбор на весь
                # файл, ни «без соединения» в строке их не снимают.
                connection_type = "flange"
                position = _with_source_text(position, item_name(text))
            elif names_no_connection(text):
                connection_type = "none"
            return with_system_column(
                original_row(
                    position, coefficients, *args, system=system, connection_type=connection_type, **kwargs
                )
            )

        pipeline.position_to_output_row = position_to_output_row


# --- Решения руководителя, которых ещё нет в коде прототипа (docs/decisions.md) ---
# Подключаются обёрткой, как правила Харнеса: прототип на сервере не меняется. Если такое же
# исправление появится в самом прототипе, обёртки ничего не изменят: купол уже задан, заголовков нет.

# Признаки размера в тексте: число из 3+ цифр, «45х45», «⌀90», «ф90», «d90», «δ=3».
_SIZE_IN_TEXT = re.compile(r"\d{3,}|\d+\s*[хxX×*]\s*\d+|[⌀∅Øøфd]\s*\d+|[δб]\s*=\s*\d", re.IGNORECASE)
# Количество в тексте строки: «2 шт», «10 мп», «кол-во: 4».
_QTY_IN_TEXT = re.compile(r"\d\s*(?:шт|м\.?п|мп|м\b|м2|м²)|кол-?во\s*[:=]?\s*\d", re.IGNORECASE)


def _number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def umbrella_with_dome(position):
    """Q38: одно число у круглого зонта — патрубок; купол = патрубок + 2 × 500 (вылет), если не указан."""

    from dataclasses import replace

    if position.element_type != "umbrella_round":
        return position
    size = dict(position.size)
    if not _number(size.get("d_neck")) and _number(size.get("d")):
        size["d_neck"] = float(size["d"])
    if not _number(size.get("d_neck")) or _number(size.get("d_dome")):
        return replace(position, size=size)
    size["d_dome"] = float(size["d_neck"]) + 1000
    questions = tuple(
        q
        for q in position.questions
        if not (
            any(key in q.code.lower() for key in ("neck", "dome"))
            or any(word in q.text.lower() for word in ("патруб", "купол", "d_dome", "d_neck"))
        )
    )
    return replace(position, size=size, questions=questions)


def is_section_heading(text: str, position) -> bool:
    """Q39: строка без количества, типа и размеров («Вытяжка из химлаборатории», шапка таблицы).

    Строка с размерами, но без количества заголовком не считается: это пропуск в заявке,
    она остаётся красной «отсутствует количество».
    """

    if position.qty is not None or position.element_type is not None or position.length_mm is not None:
        return False
    if any(_number(value) for value in position.size.values()):
        return False
    return not _SIZE_IN_TEXT.search(text) and not _QTY_IN_TEXT.search(text)


def preview_without_headings(preview, compare_completeness):
    """Убрать заголовки из разбора прототипа и заново сверить полноту (сколько позиций и количеств)."""

    from dataclasses import replace
    from types import SimpleNamespace

    # Строки, которые модель не вернула («строка не разобрана»), остаются красными: это может быть позиция.
    headings = [
        item
        for item in preview.positions
        if getattr(item, "position", None) is not None and is_section_heading(item.source_text, item.position)
    ]
    if not headings:
        return preview
    heading_nos = {str(item.source_no) for item in headings}
    positions = tuple(item for item in preview.positions if str(item.source_no) not in heading_nos)
    response = replace(
        preview.response,
        positions=tuple(p for p in preview.response.positions if str(p.source_no) not in heading_nos),
    )
    old = preview.completeness
    # Заголовок, который модель вернула как позицию, входил и в число позиций источника.
    returned = sum(1 for p in preview.response.positions if str(p.source_no) in heading_nos)
    source = SimpleNamespace(
        position_count=max(old.source_position_count - returned, 0),
        quantity_sum=old.source_quantity_sum,
    )
    completeness = compare_completeness(source, response)
    warnings = tuple(w for w in preview.warnings if w != old.message)
    if completeness.message:
        warnings += (completeness.message,)
    return replace(
        preview,
        positions=positions,
        response=response,
        completeness=completeness,
        source_position_count=completeness.source_position_count,
        model_position_count=completeness.model_position_count,
        model_quantity_sum=completeness.model_quantity_sum,
        completeness_message=completeness.message,
        warnings=warnings,
    )


def without_model_warnings(preview):
    """Убрать из предупреждений свободный текст модели разбора («element_type = null…»).

    Модель описывает свой промежуточный шаг, а не результат: особые изделия и толщины дальше считает код.
    Предупреждения кода (полнота разбора, нарезка на части) остаются.
    """

    from dataclasses import replace

    model = set(preview.response.warnings)
    if not model:
        return preview
    return replace(preview, warnings=tuple(w for w in preview.warnings if w not in model))


_COEFFICIENT_NOTE = re.compile(r"Коэффициент ([LM]) для (.+) взят из скилла")


def manager_warning(text: str) -> str:
    """Предупреждение прототипа — словами менеджера; остальные без изменений."""

    match = _COEFFICIENT_NOTE.fullmatch(text)
    if match is None:
        return text
    letter, category = match.groups()
    if (letter, category) == ("L", "нестандарт фасон"):
        return "Раскрой нестандартных изделий (зонты, клапаны, шумоглушители) — 1,5 (в шаблоне «от 1,5, уточнять»)."
    what = "Раскрой" if letter == "L" else "Коэффициент на работу"
    return f"{what} для «{category}» взят по умолчанию: в шаблоне расчётки вместо числа текст."


def _install_catalog_line_rules(pipeline) -> None:
    """ХИМВЕНТ на элементе считается. Отдельный фланец или раструб не остаётся красной строкой.
    Шибер, клапан и врезка из Excel узнаются по названию, а не по имени листа.

    Прототип на сервере не меняется. Если он уже умеет то же самое, повтор ничего не добавляет:
    покупное слово снимается один раз, строка соединения пишется один раз.
    """

    if hasattr(pipeline, "_is_scrubber_text"):
        original_scrubber = (
            getattr(pipeline, "_ameri_original_is_scrubber_text", None) or pipeline._is_scrubber_text
        )
        pipeline._ameri_original_is_scrubber_text = original_scrubber

        def _is_scrubber_text(source_text: str) -> bool:
            if is_himvent_duct(source_text):
                lower = _plain_item_text(source_text).lower()
                keywords = tuple(
                    word for word in getattr(pipeline, "PASSTHROUGH_ITEM_KEYWORDS", ()) if word != "химвент"
                )
                return any(keyword in lower for keyword in keywords)
            return original_scrubber(source_text)

        pipeline._is_scrubber_text = _is_scrubber_text

    if hasattr(pipeline, "with_special_catalog_rules"):
        original_rules = (
            getattr(pipeline, "_ameri_original_special_catalog", None) or pipeline.with_special_catalog_rules
        )
        pipeline._ameri_original_special_catalog = original_rules

        def with_special_catalog_rules(position):
            text = getattr(position, "source_text", "") or ""
            if flanged_special_kind(text) is not None:
                # Сечение — из названия, а не из имени листа и количества.
                ready = original_rules(_with_source_text(position, item_name(text)))
                return _with_source_text(ready, text)
            return original_rules(as_lone_connection(position))

        pipeline.with_special_catalog_rules = with_special_catalog_rules

    if hasattr(pipeline, "_item_head"):
        original_head = getattr(pipeline, "_ameri_original_item_head", None) or pipeline._item_head
        pipeline._ameri_original_item_head = original_head
        pipeline._item_head = lambda source_text: original_head(item_name(source_text))


def round_flange_price(d_mm, thickness_mm, material_price) -> Decimal | None:
    """Цена одного круглого фланца не из прайса: с листа 4,5 м², по наружному краю фланца (сечение + 100 мм).

    None — если с листа не выходит ни одного фланца.
    """

    sheet_price = Decimal(str(thickness_mm)) * Decimal("4.5") * Decimal("0.92") * Decimal(str(material_price))
    outer = Decimal(str(d_mm)) + 100
    per_sheet = (Decimal(7_000_000) / (outer + 159) ** 2).quantize(Decimal(1), rounding=ROUND_HALF_UP)
    if per_sheet == 0:
        return None
    return (sheet_price / per_sheet * Decimal("1.6")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def template_material_price(template: Path) -> Decimal | None:
    """Цена материала за кг из шаблона расчётки — та же ячейка, что прототип пишет в колонку J."""

    if not template.is_file():
        return None
    from openpyxl import load_workbook
    from duct_calc.xlsx_out import MATERIAL_PRICE_CELL, SHEET_NAME

    try:
        workbook = load_workbook(template, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError, zipfile.BadZipFile):
        return None
    try:
        if SHEET_NAME not in workbook.sheetnames:
            return None
        value = workbook[SHEET_NAME][MATERIAL_PRICE_CELL].value
    finally:
        workbook.close()
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    return Decimal(str(value))


def _install_flange_price(material_price: Decimal) -> None:
    """Решения руководителя по фланцам, остальное считает прототип.

    Круглый фланец не из прайса — по цене листа. Толщина фланца — только по таблице фланцев:
    толщина стенки из спецификации или от модели на фланец не переносится. Фланцы 401–449 мм — 6 мм,
    как до 400: в таблице прототипа этих размеров нет.
    """

    from duct_calc import flanges, thickness

    wall = getattr(thickness, "_ameri_original_flange_wall", None) or thickness.FLANGE_WALL
    thickness._ameri_original_flange_wall = wall
    # Границы включительные, из совпавших строк берётся большая толщина: ровно 450 мм остаётся 8 мм.
    thickness.FLANGE_WALL = (replace(wall[0], end=Decimal(450)), *wall[1:])

    original = getattr(flanges, "_ameri_original_pair_price", None) or flanges.FlangePriceBook.pair_price
    flanges._ameri_original_pair_price = original

    def pair_price(self, end, *, explicit_flange_thickness_mm=None):
        if end.kind != "round" or end.dimensions[0] in self.round_prices:
            return original(self, end)
        d = end.dimensions[0]
        one = round_flange_price(d, flanges.flange_thickness(d_mm=d), material_price)
        if one is None:
            return original(self, end)
        return flanges.PairPrice(one * 2, "calculated", f"{d.normalize():f}")

    flanges.FlangePriceBook.pair_price = pair_price


def _install_decisions(pipeline) -> None:
    """Q38 и Q39 поверх прототипа: оборачиваем его функции, исходные сохраняем для повторной установки.

    Q37 (1500 + остаток) делает сам прототип, нарезку метража не перехватываем.
    """

    neck = getattr(pipeline, "_ameri_original_umbrella_neck", None) or pipeline.with_umbrella_optional_neck_key
    pipeline._ameri_original_umbrella_neck = neck
    pipeline.with_umbrella_optional_neck_key = lambda position: umbrella_with_dome(neck(position))

    preview = getattr(pipeline, "_ameri_original_preview", None) or pipeline.preview_specification
    pipeline._ameri_original_preview = preview
    pipeline.preview_specification = lambda **kwargs: without_model_warnings(
        preview_without_headings(preview(**kwargs), pipeline.compare_completeness)
    )


def _install_harness(harness_dir: Path | None) -> None:
    """Подключить Харнес к прототипу duct-calc: правила разбора, ключевые слова, решения Q38 и Q39.

    Поверх прототипа: ХИМВЕНТ на названном элементе считается, отдельный фланец или раструб
    попадает в стоимость соединения.
    """

    from duct_calc import pipeline

    _install_decisions(pipeline)
    _install_coupling_as_socket(pipeline)
    _install_catalog_line_rules(pipeline)

    original = getattr(pipeline, "_ameri_original_prompt", None) or pipeline.build_system_prompt
    pipeline._ameri_original_prompt = original
    rules = duct_parse_rules(harness_dir)
    pipeline.build_system_prompt = (lambda: original() + "\n\n" + rules) if rules else original

    params = duct_params(harness_dir)
    for key, attr in _KEYWORD_LISTS.items():
        base = getattr(pipeline, f"_ameri_original_{attr}", None) or getattr(pipeline, attr)
        setattr(pipeline, f"_ameri_original_{attr}", base)
        extra = tuple(str(word).lower() for word in params.get(key) or ())
        setattr(pipeline, attr, tuple(base) + extra)


_RETRYABLE_HTTP_RE = re.compile(r"\bHTTP (?:5\d\d|429)\b")


def is_no_connection(error: BaseException) -> bool:
    """Ошибка клиента модели прототипа — сбой связи с DeepSeek, а не поломка разбора.

    Сетевую ошибку прототип оборачивает (`from error`), код HTTP пишет только в тексте.
    """

    return isinstance(error.__cause__, httpx.RequestError) or bool(_RETRYABLE_HTTP_RE.search(str(error)))


NO_POSITIONS_TEXT = (
    "Расчётка не построена: в файле не нашлось ни одной позиции с размерами и количеством. "
    "Уточните у клиента сечения, длины и количества изделий или спросите ассистента, что запросить."
)


def duct_calc_text(
    *,
    text: str,
    connection_label: str,
    duct_calc_dir: Path,
    base_url: str,
    api_key: str,
    model: str,
    harness_dir: Path | None = None,
) -> ActionResult:
    """Расчётка по списку позиций, введённому прямо в окне чата."""

    payload = text.strip()
    if not payload:
        return ActionResult(
            "Для расчётки приложите файл спецификации или напишите список позиций в сообщении."
        )
    with tempfile.TemporaryDirectory() as tmp:
        source_name = "позиции-из-чата.md"
        source = Path(tmp) / source_name
        source.write_text(payload + "\n", encoding="utf-8")
        result = duct_calc(
            spec=Attachment(id=0, name=source_name, path=source, size=source.stat().st_size),
            connection_label=connection_label,
            duct_calc_dir=duct_calc_dir,
            base_url=base_url,
            api_key=api_key,
            model=model,
            harness_dir=harness_dir,
        )
    prefix = f"Расчётка по файлу «{source_name}»"
    if result.text.startswith(prefix):
        return ActionResult(text=result.text.replace(prefix, "Расчётка по тексту из чата", 1), files=result.files)
    return result


def duct_calc(
    *,
    spec: Attachment,
    connection_label: str,
    duct_calc_dir: Path,
    base_url: str,
    api_key: str,
    model: str,
    harness_dir: Path | None = None,
) -> ActionResult:
    """Расчётка воздуховодов по вложенной Спецификации (прототип duct-calc)."""

    problem = key_problem(api_key)
    if problem:
        return ActionResult(text=problem)
    _import_duct_calc(duct_calc_dir)
    _install_harness(harness_dir)
    from duct_calc.model_client import DeepSeekClient, ModelClientConfig, ModelClientError
    from duct_calc.pipeline import PipelineNeedsInput, run_specification_pipeline

    client = DeepSeekClient(
        ModelClientConfig(
            base_url=base_url,
            api_key=api_key,
            model=model,
            # Прототип передаёт значение в httpx как есть, поэтому соединение и ответ ограничены раздельно.
            # Ограничение — на каждую часть спецификации (~90 строк); один повтор при сбое делает прототип.
            timeout_seconds=DEEPSEEK_TIMEOUT,
        )
    )
    with tempfile.TemporaryDirectory() as tmp:
        output = Path(tmp) / f"расчетка_{spec.path.stem}.xlsx"
        template = duct_calc_dir / "data" / DUCT_TEMPLATE_NAME
        material_price = template_material_price(template)
        if material_price is not None:
            # Без числа в шаблоне прототип сам остановится с понятной ошибкой.
            _install_flange_price(material_price)
        connection = CONNECTION_TYPES[connection_label]

        def run_lines() -> tuple[object, str]:
            # Таблицы .md/.odt/.csv/.xlsx без «№» прототип разбирает хуже, чем строки заказа Word: передаём ему
            # спецификацию построчно «Наименование — количество ед.», без шапки и номеров строк.
            source = Path(tmp) / f"{spec.path.stem}.md"
            source.write_text(spec_lines_text(spec.path) or "", encoding="utf-8")
            result = run_specification_pipeline(
                source_path=source,
                output_path=output,
                connection_type=connection,
                template_path=template,
                model_client=client,
            )
            return result, f"Позиций: {result.completeness.model_position_count}."

        suffix = spec.path.suffix.lower()
        try:
            if suffix == ".csv":
                # Сначала CSV-модуль прототипа (свой формат выгрузки); обычную таблицу он не читает.
                from duct_calc.csv_pipeline import run_csv_pipeline

                try:
                    result = run_csv_pipeline(
                        source_path=spec.path,
                        output_path=output,
                        connection_type=connection,
                        template_path=template,
                        model_client=client,
                    )
                    summary = f"Позиций: {result.source_item_count}, строк в расчётке: {result.output_row_count}."
                except (ValueError, KeyError) as error:
                    if isinstance(error, PipelineNeedsInput):
                        raise
                    result, summary = run_lines()
            elif suffix in {".md", ".odt"} or (suffix == ".xlsx" and spec_lines_text(spec.path) is not None):
                result, summary = run_lines()
            else:
                result = run_specification_pipeline(
                    source_path=spec.path,
                    output_path=output,
                    connection_type=connection,
                    template_path=template,
                    model_client=client,
                )
                summary = f"Позиций: {result.completeness.model_position_count}."
        except PipelineNeedsInput as error:
            return ActionResult(text=f"Нужно уточнение, расчётка не построена: {error}")
        except ModelClientError as error:
            if not is_no_connection(error):
                raise
            return ActionResult(text=NO_CONNECTION_MESSAGE)
        except ValueError as error:
            if "без позиций" not in str(error):
                raise
            # Прототип не строит пустую расчётку: в файле нет ни одной позиции с размером или количеством.
            return ActionResult(text=NO_POSITIONS_TEXT)
        lines = [f"Расчётка по файлу «{spec.name}» готова ({connection_label.lower()}). {summary}"]
        red = red_positions(output)
        if red:
            lines.append(
                "Красные позиции — не посчитаны, нужно уточнить:\n" + "\n".join(f"- {item}" for item in red)
            )
        lines += result_notes(output, connection_label)
        notes = [manager_warning(w) for w in result.warnings if _COEFFICIENT_NOTE.fullmatch(w)]
        warnings = [w for w in result.warnings if not _COEFFICIENT_NOTE.fullmatch(w)]
        lines += notes
        if warnings:
            lines.append("Предупреждения: " + "; ".join(warnings))
        return ActionResult(text="\n\n".join(lines), files=[(output.name, output.read_bytes())])
